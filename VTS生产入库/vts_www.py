import requests
import json
import time
import openpyxl
from openpyxl import load_workbook
from datetime import datetime
import re

# ====================== 配置区 ======================
BASE_URL = "http://10.224.245.101:8082/Packing/PackCarton.aspx"
EXCEL_PATH = "date_www.xlsx"
LOG_PATH   = "log_www.xlsx"
XUSER = "N33274"
IP_STATION = "10.224.151.59"
SLEEP_SEC = 1
# ====================================================

session = requests.Session()

LOG_HEADER = ["时间","工单","箱号","箱容","起始SN","结束SN","实际装箱数","是否满箱","操作结果"]

def init_log_file():
    try:
        load_workbook(LOG_PATH)
    except FileNotFoundError:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "箱操作日志"
        ws.append(LOG_HEADER)
        wb.save(LOG_PATH)

def write_box_log(lot, carton, carton_size, start_sn, end_sn, qty, full_flag, result):
    wb = load_workbook(LOG_PATH)
    ws = wb["箱操作日志"]
    row = [
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        lot, carton, carton_size, start_sn, end_sn, qty,
        "是" if full_flag else "否",
        result
    ]
    ws.append(row)
    wb.save(LOG_PATH)
    print(f"[日志已写入] 箱:{carton} 数量:{qty} 满箱:{full_flag} 结果:{result}")


def get_pack_list_bylot(lot_name:str):
    params = {
        "Oper":"get_pack_list_bylot",
        "LotName": lot_name,
        "_": int(time.time()*1000)
    }
    resp = session.get(BASE_URL, params=params, timeout=15)
    return resp.json()["msg"]


def filter_target_carton(carton_list):
    """筛选未满箱子：QTY < CARTON_SIZE，满箱直接丢弃"""
    result = []
    for item in carton_list:
        qty = int(item["QTY"])
        carton_size = int(item["CARTON_SIZE"])
        if qty < carton_size:
            result.append(item)
    return result


def take_over_old_carton(lot_name:str):
    """接管旧箱号，Is_new=0，实现网页select_old_carton解锁逻辑"""
    params = {
        "Oper":"get_pk_rule",
        "LotName": lot_name,
        "Is_new":0,
        "IP": IP_STATION,
        "_": int(time.time()*1000)
    }
    resp = session.get(BASE_URL, params=params, timeout=15)
    print(f"[接管旧箱] get_pk_rule(Is_new=0) 返回:{resp.text}")
    time.sleep(SLEEP_SEC)


def save_pack_carton(sn:str, station:str, lot_name:str, carton:str):
    params = {
        "Oper":"save_pack_carton",
        "StepName":"",
        "SN": sn,
        "Station": station,
        "LotName": lot_name,
        "IP": IP_STATION,
        "CusCarton":"",
        "Carton": carton,
        "Xuser": XUSER,
        "_": int(time.time()*1000)
    }
    resp = session.get(BASE_URL, params=params, timeout=15)
    print(f"[写入SN] SN={sn}, Carton={carton}, 返回:{resp.text}")
    time.sleep(SLEEP_SEC)
    return resp.text


def get_cart_contents(carton_sn:str):
    params = {
        "Oper":"get_pack_list_bycarton",
        "Carton": carton_sn,
        "_": int(time.time()*1000)
    }
    resp = session.get(BASE_URL, params=params, timeout=15)
    item_list = resp.json()["msg"]
    current_count = len(item_list)
    return current_count, item_list


def get_pk_rule_new(lot_name:str):
    """满箱收尾 Is_new=1"""
    params = {
        "Oper":"get_pk_rule",
        "LotName": lot_name,
        "Is_new":1,
        "IP": IP_STATION,
        "_": int(time.time()*1000)
    }
    resp = session.get(BASE_URL, params=params, timeout=15)
    print(f"[满箱收尾get_pk_rule]返回:{resp.text}")
    return resp.text


def generate_sn_range(start_sn:str, end_sn:str):
    """
    【新版安全拆分】自动分割文字前缀+尾部数字
    示例: "2635018705B00121" →前缀:"2635018705B" 起始数字:121
    """
    # 正则分割：前面非数字 + 后面数字
    match_start = re.match(r'(.*?)(\d+)$', start_sn)
    match_end   = re.match(r'(.*?)(\d+)$', end_sn)
    if not match_start or not match_end:
        raise Exception("SN解析失败，无法分割前缀与尾部数字")

    prefix = match_start.group(1)
    start_num = int(match_start.group(2))
    end_num = int(match_end.group(2))
    # 获取尾部原始补零长度
    tail_len = len(match_start.group(2))

    sn_list = []
    for num in range(start_num, end_num + 1):
        sn = f"{prefix}{num:0{tail_len}d}"
        sn_list.append(sn)
    return sn_list


def read_excel(filepath):
    wb = load_workbook(filepath)
    ws = wb.active
    rows = []
    for r in range(2, ws.max_row+1):
        so      = str(ws.cell(r,1).value).strip()
        station = str(ws.cell(r,2).value).strip()
        start   = str(ws.cell(r,3).value).strip()
        end     = str(ws.cell(r,4).value).strip()
        rows.append({"SO":so,"STATION":station,"START_SN":start,"END_SN":end})
    return rows


def main():
    init_log_file()
    excel_rows = read_excel(EXCEL_PATH)
    if len(excel_rows) == 0:
        print("Excel没有数据")
        return

    for task in excel_rows:
        lot_name  = task["SO"]
        station   = task["STATION"]
        sn_pool   = generate_sn_range(task["START_SN"], task["END_SN"])
        total_sn  = len(sn_pool)
        sn_ptr    = 0

        print(f"\n==== 工单 {lot_name} 待处理SN总数:{total_sn} ====")

        # 外层循环，一箱装完刷新箱子列表继续下一箱
        while sn_ptr < total_sn:
            carton_all = get_pack_list_bylot(lot_name)
            target_cartons = filter_target_carton(carton_all)

            if not target_cartons:
                print("【警告】没有剩余未满箱子，工单装箱结束！")
                break

            carton_info = target_cartons[0]
            carton_sn = carton_info["PACK_PARENT_NUMBER"]
            target_full_qty = int(carton_info["CARTON_SIZE"])
            is_ok_flag = carton_info["IS_OK"]
            box_start_index = sn_ptr
            box_start_sn = sn_pool[sn_ptr]

            try:
                # 如果箱子状态OK，调用接管接口，替代失败的unlink_ip解绑
                if is_ok_flag == "OK":
                    take_over_old_carton(lot_name)

                box_full = False
                while True:
                    current_num,_ = get_cart_contents(carton_sn)
                    print(f"箱子 {carton_sn} 当前:{current_num} / 需装满:{target_full_qty}")

                    if current_num >= target_full_qty:
                        get_pk_rule_new(lot_name)
                        box_full = True
                        break

                    if sn_ptr >= total_sn:
                        box_full = False
                        break

                    current_sn = sn_pool[sn_ptr]
                    resp_text = save_pack_carton(current_sn, station, lot_name, carton_sn)

                    if "already packed" in resp_text:
                        print(f"跳过已装箱SN:{current_sn}")
                        sn_ptr += 1
                        continue
                    if "Serial number not exists" in resp_text:
                        print(f"【告警】SN不存在跳过:{current_sn}，继续下一个SN")
                        sn_ptr += 1
                        continue

                    sn_ptr += 1

                box_end_sn = sn_pool[sn_ptr - 1] if sn_ptr>0 else ""
                write_box_log(lot_name, carton_sn, target_full_qty,
                              box_start_sn, box_end_sn, current_num,
                              box_full, "成功")

            except Exception as e:
                write_box_log(lot_name, carton_sn, target_full_qty,
                              box_start_sn, "", -1, False, f"失败:{str(e)}")
                print(f"[异常]箱 {carton_sn} 错误:{e}")
                break

if __name__ == "__main__":
    main()
