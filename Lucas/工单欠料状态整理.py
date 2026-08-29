import sys
import subprocess
import importlib.util
from datetime import datetime

# ===================== 自动安装依赖 =====================
def auto_install(package_name: str):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package_name])

for pkg in ["requests", "openpyxl", "pyperclip"]:
    if not importlib.util.find_spec(pkg):
        print(f"[自动安装] {pkg} 未找到，正在安装...")
        auto_install(pkg)

import hashlib
import requests
import pyperclip
from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from typing import Optional, List, Dict, Any

# ===================== 全局配置 =====================
GLOBAL_COMPANY_ID = "6801"
USER_NAME = "1441"
PASSWORD = "1442"

# ===================== MES 登录模块 =====================
class MES:
    def encrypt_password(self, password: str) -> str:
        md5 = hashlib.md5()
        md5.update(password.encode('utf-8'))
        return md5.hexdigest()

    def login(self, account: str, password: str) -> Optional[str]:
        print("[MES] 开始登录...")
        encrypted_pwd = self.encrypt_password(password)
        url = "http://10.97.245.205:86/login"
        headers = {
            "Content-Type": "application/json;charset=UTF-8",
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
        }
        payload = {"account": account, "password": encrypted_pwd}
        try:
            resp = requests.post(url, headers=headers, json=payload, timeout=15)
            resp.raise_for_status()
            res_json = resp.json()
            if res_json.get("code") == 200 and "data" in res_json and "token" in res_json["data"]:
                print("[MES] 登录成功，已获取 Token")
                return res_json["data"]["token"]
            else:
                print(f"[MES] 登录失败: {res_json.get('info', '未知错误')}")
                return None
        except Exception as e:
            print(f"[MES] 登录异常: {str(e)}")
            return None

# ===================== 库存下载函数 =====================
def fetch_inventory_page(token: str, company_id: str, page: int, rows: int = 10000) -> Dict[str, Any]:
    url = "http://10.97.245.205:92/prod_baseapi/wms/sendproductionbill/getSock/page"
    params = {
        "rows": rows,
        "page": page,
        "sidx": "part_no DESC, vtech_lot ASC",
        "part_no": "",
        "bin": "",
        "company_id": company_id,
        "statusList": "",
        "batch": "",
        "mgl": "false",
        "uskey": "false",
        "RMBPurchaseKey": "false"
    }
    headers = {
        "token": token,
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
    }
    try:
        resp = requests.get(url, params=params, headers=headers, timeout=30)
        resp.raise_for_status()
        data = resp.json()
        if data.get("code") == 200:
            return data["data"]
        else:
            print(f"[错误] API 返回 code {data.get('code')}: {data.get('info')}")
            return {"rows": [], "total": 0}
    except Exception as e:
        print(f"[错误] 获取第 {page} 页失败: {e}")
        return {"rows": [], "total": 0}

def fetch_all_inventory(token: str, company_id: str) -> List[Dict]:
    print("[信息] 开始下载全部库存数据...")
    rows_per_page = 10000
    result = fetch_inventory_page(token, company_id, 1, rows_per_page)
    if result.get("total", 0) > 0 or result.get("rows"):
        total = result.get("total", 0)
        if total <= rows_per_page:
            print(f"[信息] 一次性获取成功，共 {len(result['rows'])} 条记录")
            return result["rows"]
        else:
            print(f"[信息] 总记录数 {total} 超过单页限制，启动分页获取...")
            all_rows = result["rows"]
            page = 2
            while len(all_rows) < total:
                next_page = fetch_inventory_page(token, company_id, page, rows_per_page)
                rows = next_page.get("rows", [])
                if not rows:
                    break
                all_rows.extend(rows)
                print(f"[信息] 已获取 {len(all_rows)} / {total} 条")
                page += 1
            print(f"[信息] 下载完成，共 {len(all_rows)} 条记录")
            return all_rows
    else:
        print("[警告] 一次性获取失败，将使用分页模式（每页 50 条）")
        all_rows = []
        page = 1
        rows_per_page = 50
        total = None
        while True:
            result = fetch_inventory_page(token, company_id, page, rows_per_page)
            page_rows = result.get("rows", [])
            if not page_rows:
                break
            all_rows.extend(page_rows)
            total = result.get("total", 0)
            print(f"[信息] 已获取第 {page} 页，累计 {len(all_rows)} / {total} 条")
            if len(all_rows) >= total:
                break
            page += 1
        print(f"[信息] 下载完成，共 {len(all_rows)} 条记录")
        return all_rows

# ===================== 智能解析剪贴板数据 =====================
def parse_clipboard_data(text: str) -> List[Dict]:
    """自动检测是否有表头，若无则按固定列位置解析（第4列为Component）"""
    lines = text.strip().splitlines()
    if not lines:
        return []

    first_line = lines[0].strip()
    has_header = False
    if '\t' in first_line:
        parts = first_line.split('\t')
    else:
        parts = first_line.split()
    for p in parts:
        if p.strip().lower() == "component":
            has_header = True
            break

    if has_header:
        header = [h.strip() for h in parts]
        rows = []
        for line in lines[1:]:
            if '\t' in line:
                vals = line.split('\t')
            else:
                vals = line.split()
            while len(vals) < len(header):
                vals.append("")
            row_dict = dict(zip(header, vals))
            rows.append(row_dict)
        return rows
    else:
        default_headers = ["col1", "col2", "col3", "Component", "col5", "col6",
                           "col7", "col8", "col9", "col10", "col11", "col12"]
        rows = []
        for line in lines:
            if '\t' in line:
                vals = line.split('\t')
            else:
                vals = line.split()
            while len(vals) < 4:
                vals.append("")
            row_dict = {}
            for idx, h in enumerate(default_headers):
                if idx < len(vals):
                    row_dict[h] = vals[idx]
                else:
                    row_dict[h] = ""
            rows.append(row_dict)
        return rows

# ===================== 构建库存查找表 =====================
def build_inventory_lookup(inventory_rows: List[Dict]) -> Dict[str, List[Dict]]:
    lookup = {}
    for item in inventory_rows:
        part_no = item.get("part_no", "").strip()
        if not part_no:
            continue
        lookup.setdefault(part_no, []).append(item)
    return lookup

# ===================== 格式化数量（去掉 .0） =====================
def format_qty(val):
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    if isinstance(val, str):
        try:
            f = float(val)
            if f.is_integer():
                return str(int(f))
            else:
                return val.rstrip('0').rstrip('.') if '.' in val else val
        except:
            return val
    return str(val)

# ===================== 生成 Excel =====================
def generate_excel(processed_rows: List[Dict], output_file: str = "Shortage_of_FG.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "匹配结果"

    new_cols = ["Customer", "RT", "Location", "Stock_Info", "Special_Stock", "Dept"]
    if processed_rows:
        sample_keys = [k for k in processed_rows[0].keys() if k not in new_cols]
    else:
        sample_keys = []

    friendly_headers = ["Order", "Material", "Plant", "Component", "Req. Qty",
                        "Withdrawn Qty", "Assign Stock", "Assign Loc", "Batch",
                        "Unit", "WBS Element", "Message"]

    if sample_keys and sample_keys[0].startswith("col"):
        original_headers = friendly_headers[:len(sample_keys)]
    else:
        original_headers = sample_keys

    headers = original_headers + new_cols
    ws.append(headers)

    numeric_columns = {"Req. Qty", "Withdrawn Qty", "Assign Stock"}

    for row in processed_rows:
        row_data = []
        for h in original_headers:
            val = row.get(h, "")
            if val == "" and h in friendly_headers:
                idx = friendly_headers.index(h) if h in friendly_headers else -1
                if idx >= 0:
                    col_key = f"col{idx+1}"
                    val = row.get(col_key, "")
            if h in numeric_columns and isinstance(val, str):
                clean = val.replace(',', '').replace(' ', '').strip()
                if clean == '':
                    val = None
                else:
                    try:
                        if '.' in clean:
                            val = float(clean)
                        else:
                            val = int(clean)
                    except ValueError:
                        pass
            row_data.append(val)
        for new_h in new_cols:
            row_data.append(row.get(new_h, ""))
        ws.append(row_data)

    # 设置列宽和自动换行
    for col_idx, col_name in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 20
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=col_idx).alignment = Alignment(wrap_text=True)

    # ===== 行颜色标记（根据“Location”列） =====
    pos_col_idx = None
    for idx, h in enumerate(headers, start=1):
        if h == "Location":
            pos_col_idx = idx
            break
    if pos_col_idx is None:
        pos_col_idx = len(original_headers) + 3

    green_fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")
    light_red_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")

    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=pos_col_idx)
        pos_text = cell.value or ""
        if not pos_text:
            continue
        has_qp00 = "QP00" in pos_text
        has_q_other = "Q" in pos_text and not has_qp00
        if has_qp00:
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col).fill = green_fill
        elif has_q_other:
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col).fill = light_red_fill

    # ===== 设置“Special_Stock”列中值为"K"的单元格背景为黄色 =====
    special_col_idx = None
    for idx, h in enumerate(headers, start=1):
        if h == "Special_Stock":
            special_col_idx = idx
            break
    if special_col_idx:
        yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=special_col_idx)
            if cell.value == "K":
                cell.fill = yellow_fill

    wb.save(output_file)
    print(f"[信息] Excel 已保存至: {output_file}")

# ===================== 主流程 =====================
def main():
    mes = MES()
    token = mes.login(USER_NAME, PASSWORD)
    if not token:
        print("[错误] 登录失败，程序退出。")
        return

    inventory_rows = fetch_all_inventory(token, GLOBAL_COMPANY_ID)
    if not inventory_rows:
        print("[错误] 未获取到任何库存数据。")
        return

    lookup = build_inventory_lookup(inventory_rows)
    print(f"[信息] 库存查找表构建完成，共 {len(lookup)} 个不同的物料号。")

    sample_parts = list(lookup.keys())[:5]
    print("[调试] 库存中的前5个 part_no 示例:")
    for p in sample_parts:
        print(f"    {p}")

    clipboard_text = pyperclip.paste()
    if not clipboard_text:
        print("[错误] 剪贴板为空。")
        return
    clipboard_rows = parse_clipboard_data(clipboard_text)
    if not clipboard_rows:
        print("[错误] 剪贴板数据解析失败。")
        return

    print(f"[信息] 从剪贴板解析到 {len(clipboard_rows)} 行数据。")

    # ======== 逐行处理（不分组） ========
    result_rows = []
    unmatched_examples = []

    for row in clipboard_rows:
        # 初始化新增列
        row["Customer"] = ""
        row["RT"] = ""
        row["Location"] = ""
        row["Stock_Info"] = ""
        row["Special_Stock"] = ""
        row["Dept"] = "PMC"   # 默认

        component = row.get("Component", "").strip()
        if not component:
            # 无 Component，保留空新增列，Dept 默认 PMC
            result_rows.append(row)
            continue

        matched_items = lookup.get(component, [])
        if not matched_items:
            row["Stock_Info"] = "No stock found"
            if len(unmatched_examples) < 5 and component not in unmatched_examples:
                unmatched_examples.append(component)
            result_rows.append(row)
            continue

        # 去重库存记录
        seen = set()
        unique_items = []
        for item in matched_items:
            key = (item.get("wbs",""), item.get("vtech_lot",""), item.get("bin_id",""))
            if key not in seen:
                seen.add(key)
                unique_items.append(item)

        # 填充各列
        first_wbs = unique_items[0].get("wbs", "") or ""
        rt_list = []
        bin_list = []
        info_list = []
        for item in unique_items:
            rt_list.append(item.get("vtech_lot", "") or "")
            bin_list.append(item.get("bin_id", "") or "")
            store_qty = item.get("store_qty", 0)
            q_qty = item.get("q_qty", 0)
            try:
                store_qty = float(store_qty) if store_qty not in (None, "") else 0
            except:
                store_qty = 0
            try:
                q_qty = float(q_qty) if q_qty not in (None, "") else 0
            except:
                q_qty = 0
            qty_parts = []
            if store_qty != 0:
                qty_parts.append(f"Stock:{format_qty(store_qty)}")
            if q_qty != 0:
                qty_parts.append(f"Pending:{format_qty(q_qty)}")
            if not qty_parts:
                qty_parts.append("No quantity")
            info_list.append(", ".join(qty_parts))

        row["Customer"] = first_wbs
        row["RT"] = "\n".join(rt_list)
        row["Location"] = "\n".join(bin_list)
        row["Stock_Info"] = "\n".join(info_list)

        # Special_Stock 取第一条记录的 special_stock
        special_stock = unique_items[0].get("special_stock", "") or ""
        row["Special_Stock"] = special_stock

        # 根据 Location 设置 Dept
        pos_text = row["Location"]
        if "QP00" in pos_text:
            row["Dept"] = "Warehouse"
        elif "Q" in pos_text and "QP00" not in pos_text:
            row["Dept"] = "QA"
        else:
            row["Dept"] = "PMC"

        result_rows.append(row)

    # ======== 按 Component 排序 ========
    result_rows.sort(key=lambda x: (x.get("Component", "") or "zzzzzzzzz"))

    print(f"[信息] 共处理 {len(result_rows)} 行。")
    if unmatched_examples:
        print("[提示] 以下 Component 未匹配到库存（仅显示前5个）:")
        for ex in unmatched_examples:
            print(f"    - {ex}")
        print("    请对比上述 Component 与库存 part_no 的格式是否完全一致。")

    generate_excel(result_rows)
    print("[完成] 处理结束。")

if __name__ == "__main__":
    main()