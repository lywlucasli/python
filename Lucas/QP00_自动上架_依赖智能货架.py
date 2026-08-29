# -*- coding: utf-8 -*-
"""
MES 自动上架工具 - 交互式分步执行版（强制使用 SA00 库位）
包含日志记录和 Excel 结果报表，详细反馈每个 RT 的执行结果。
"""

import hashlib
import requests
import time
import json
from typing import Optional, Dict, List, Tuple
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import os
import sys
import logging
from datetime import datetime

# ======================== 配置（请按实际修改）========================
MES_LOGIN_URL = "http://10.97.245.205:86/login"
MES_ACCOUNT = "NAF684"          # 您的登录账号
MES_PASSWORD = "666888"         # 您的登录密码

# 待上架报表 API
WAITING_REPORT_URL = "http://10.97.245.205:92/prod_baseapi/wms/instorereport/getPendingRep/page"
COMPANY_ID = "6801"              # 公司 ID，固定

# 智能料架 API
SMART_SHELF_URL = "http://192.168.9.100:8081/mate/searchByInfo"
SMART_USER_NAME = "NAF684"
SMART_USER_GROUP_POWER = "1"
SMART_USER_LANGUAGE = "zh"

# 上架相关 API
GET_RT_INFO_URL = "http://10.97.245.205:92/prod_baseapi/wms/sendproductionbill/getRTInfo"
PUTAWAY_URL = "http://10.97.245.205:92/prod_baseapi/CommonAPI/CallSp"

# 输出文件
WAITING_REPORT_EXCEL = "waiting_report.xlsx"
WAITING_REPORT_JSON = "waiting_report.json"
SMART_SHELF_JSON = "smart_shelf_stats.json"

# 重试配置
MAX_RETRY = 3
RETRY_DELAY = 2

# ========== 强制使用的库位（所有上架均使用此库位） ==========
FORCED_LOCATION = "SA00"   # 您可以修改此值，但必须确保是允许的库位

# ========== 日志目录 ==========
LOG_DIR = "logs"
os.makedirs(LOG_DIR, exist_ok=True)

# ======================== 日志配置 ================================
def setup_logger():
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_filename = os.path.join(LOG_DIR, f"auto_putaway_{timestamp}.log")
    logger = logging.getLogger("MES_auto_putaway")
    logger.setLevel(logging.INFO)
    if logger.handlers:
        return logger
    fh = logging.FileHandler(log_filename, encoding='utf-8')
    fh.setLevel(logging.INFO)
    ch = logging.StreamHandler(sys.stdout)
    ch.setLevel(logging.INFO)
    formatter = logging.Formatter('%(asctime)s - %(levelname)s - %(message)s')
    fh.setFormatter(formatter)
    ch.setFormatter(formatter)
    logger.addHandler(fh)
    logger.addHandler(ch)
    logger.info(f"日志文件: {log_filename}")
    return logger

# 全局日志对象
logger = None

# ======================== 核心类 ============================

class MESAuth:
    @staticmethod
    def encrypt_password(password: str) -> str:
        md5 = hashlib.md5()
        md5.update(password.encode('utf-8'))
        return md5.hexdigest()

    @staticmethod
    def login(account: str, password: str) -> Optional[str]:
        global logger
        logger.info("正在登录 MES...")
        encrypted_pwd = MESAuth.encrypt_password(password)
        headers = {
            "Accept": "application/json, text/plain, */*",
            "Accept-Encoding": "gzip, deflate",
            "Accept-Language": "zh-CN,zh;q=0.9",
            "Connection": "keep-alive",
            "Content-Type": "application/json;charset=UTF-8",
            "Host": "10.97.245.205:86",
            "Origin": "http://10.97.245.205:92",
            "Referer": "http://10.97.245.205:92/",
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
        }
        data = {"account": account, "password": encrypted_pwd}
        try:
            resp = requests.post(MES_LOGIN_URL, headers=headers, json=data, timeout=10)
            resp.raise_for_status()
            result = resp.json()
            if result.get("code") == 200 and "data" in result and "token" in result["data"]:
                token = result["data"]["token"]
                logger.info("登录成功，获取 Token")
                return token
            else:
                logger.error(f"登录失败: {result.get('info', '未知错误')}")
                return None
        except Exception as e:
            logger.exception(f"登录异常: {e}")
            return None


class SmartShelfDownloader:
    @staticmethod
    def build_payload(start: int, rows: int) -> dict:
        return {
            "part_num": "",
            "warehouse_category_id": 0,
            "warehouse_name_id": 0,
            "save_id": "",
            "lot_code": "",
            "mfg_date_start": "",
            "supplier_name": "",
            "mfg_date_end": "",
            "start": start,
            "rows": rows,
            "user_name": SMART_USER_NAME,
            "user_group_power": SMART_USER_GROUP_POWER,
            "user_language": SMART_USER_LANGUAGE,
            "shelf_id": "",
            "position_info": "",
            "start_date": "",
            "end_date": "",
        }

    @staticmethod
    def fetch_page(start: int, rows: int = 10000) -> Optional[dict]:
        global logger
        payload = SmartShelfDownloader.build_payload(start, rows)
        for attempt in range(1, MAX_RETRY + 1):
            try:
                resp = requests.post(SMART_SHELF_URL, json=payload, timeout=30)
                resp.raise_for_status()
                data = resp.json()
                if data.get("result") == 1:
                    return data
                else:
                    logger.warning(f"智能料架第{attempt}次 result!=1: {data}")
            except Exception as e:
                logger.warning(f"智能料架第{attempt}次请求异常: {e}")
            if attempt < MAX_RETRY:
                time.sleep(RETRY_DELAY)
        return None

    @staticmethod
    def extract_rt_from_label(label: str) -> str:
        if not label:
            return ""
        try:
            parts = label.split('#')
            if len(parts) >= 2:
                return parts[1].strip()
            return ""
        except:
            return ""

    @staticmethod
    def download_and_stats() -> Dict[str, float]:
        global logger
        rt_quantity_sum = {}
        start = 0
        consecutive_empty = 0
        total = None
        logger.info("开始下载智能货架数据...")
        while True:
            data = SmartShelfDownloader.fetch_page(start)
            if data is None:
                logger.error(f"start={start} 重试失败，停止")
                break

            mate_list = data.get("mate_list", [])
            if total is None:
                total = data.get("rows", 0)
                logger.info(f"智能货架声明总记录数: {total}")

            if not mate_list:
                consecutive_empty += 1
                logger.info(f"  start={start} 返回空页 ({consecutive_empty}/3)")
                if consecutive_empty >= 3:
                    logger.info("连续3次空页，数据拉取完成")
                    break
                start += 10000
                time.sleep(0.5)
                continue

            consecutive_empty = 0
            for item in mate_list:
                label = item.get("label_code", "")
                rt = SmartShelfDownloader.extract_rt_from_label(label)
                if rt:
                    qty_str = item.get("quantity", "0")
                    try:
                        qty = float(qty_str) if qty_str else 0.0
                    except:
                        qty = 0.0
                    rt_quantity_sum[rt] = rt_quantity_sum.get(rt, 0.0) + qty

            logger.info(f"已获取 {len(mate_list)} 条 (累计统计到 {len(rt_quantity_sum)} 个 RT)")
            if total and len(mate_list) < 10000:
                logger.info("本页数量小于请求行数，已到最后一页")
                break
            start += 10000
            time.sleep(0.1)

        logger.info(f"智能货架统计完成，共 {len(rt_quantity_sum)} 个 RT")
        return rt_quantity_sum

    @staticmethod
    def save_stats_to_json(stats: Dict[str, float], filename: str):
        with open(filename, 'w', encoding='utf-8') as f:
            json.dump(stats, f, ensure_ascii=False, indent=2)
        logger.info(f"智能货架统计已保存至 {filename}")


class WaitingReportDownloader:
    @staticmethod
    def fetch_page(token: str, page: int, rows: int = 100000) -> Optional[dict]:
        global logger
        headers = {
            "Accept": "application/json, text/plain, */*",
            "Accept-Encoding": "gzip, deflate",
            "Accept-Language": "zh-CN,zh;q=0.9",
            "Connection": "keep-alive",
            "Host": "10.97.245.205:92",
            "Referer": "http://10.97.245.205:92/wms/report/PendingRep/index",
            "token": token,
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
        }
        params = {
            "rows": rows,
            "page": page,
            "sidx": "vtech_lot",
            "sord": "asc",
            "start_date": "",
            "end_date": "",
            "company_id": COMPANY_ID,
            "startTime": "00:00:00",
            "endTime": "00:00:00",
            "partNo": "",
            "batch": ""
        }
        for attempt in range(1, MAX_RETRY + 1):
            try:
                resp = requests.get(WAITING_REPORT_URL, headers=headers, params=params, timeout=30)
                resp.raise_for_status()
                data = resp.json()
                if data.get("code") == 200:
                    return data.get("data", {})
                else:
                    logger.warning(f"第{attempt}次请求返回 code!=200: {data}")
            except Exception as e:
                logger.warning(f"第{attempt}次请求异常: {e}")
            if attempt < MAX_RETRY:
                time.sleep(RETRY_DELAY)
        return None

    @staticmethod
    def download_all(token: str) -> List[dict]:
        global logger
        all_rows = []
        page = 1
        consecutive_empty = 0
        while True:
            logger.info(f"获取待上架报表第 {page} 页...")
            data = WaitingReportDownloader.fetch_page(token, page)
            if data is None:
                logger.warning(f"第 {page} 页获取失败，跳过")
                consecutive_empty += 1
                if consecutive_empty >= 3:
                    logger.warning("连续 3 次失败，终止")
                    break
                page += 1
                continue

            rows = data.get("rows", [])
            records = data.get("records", 0)
            logger.info(f"  返回 {len(rows)} 条，总记录数 {records}")

            if not rows:
                consecutive_empty += 1
                if consecutive_empty >= 3:
                    logger.info("连续 3 次空页，认为数据已全部拉取")
                    break
            else:
                consecutive_empty = 0
                all_rows.extend(rows)

            if len(rows) < 100000:
                logger.info("本页数量小于请求行数，已到最后一页")
                break
            if records > 0 and len(all_rows) >= records:
                logger.info(f"已获取 {len(all_rows)} 条，达到 records 总数")
                break

            page += 1
            time.sleep(0.3)

        return all_rows

    @staticmethod
    def save_to_excel(rows: List[dict], filename: str):
        global logger
        headers = [
            "序号", "Part No", "RT No", "Warehouse", "WBS", "Pass Qty",
            "Location", "COO", "PN Description", "Vendor Name", "DN No.",
            "PO", "PO Line", "Packing No.", "Type", "Approved Date", "vtech_lot"
        ]
        wb = Workbook()
        ws = wb.active
        ws.title = "待上架报表"
        ws.append(headers)
        for item in rows:
            row = [
                item.get("rowindex", ""),
                item.get("part_no", ""),
                item.get("vtech_lot", ""),
                item.get("wh_id", ""),
                item.get("wbs", ""),
                item.get("unit_qty", ""),
                item.get("location", ""),
                item.get("coo", ""),
                item.get("part_desc", ""),
                item.get("vendor_name", ""),
                item.get("invoice_no", ""),
                item.get("po_no", ""),
                item.get("po_item", ""),
                item.get("plate_number", ""),
                "",
                item.get("credit_date", ""),
                item.get("vtech_lot", "")
            ]
            ws.append(row)
        wb.save(filename)
        logger.info(f"待上架报表已保存至 {filename}")

    @staticmethod
    def extract_waiting_info(rows: List[dict]) -> Dict[str, dict]:
        info = {}
        for item in rows:
            vtech = item.get("vtech_lot", "")
            if not vtech:
                continue
            qty = item.get("unit_qty", 0.0)
            try:
                qty = float(qty)
            except:
                qty = 0.0
            if qty <= 0:
                continue
            if vtech not in info:
                info[vtech] = {"total_qty": 0.0}
            info[vtech]["total_qty"] += qty
        return info

    @staticmethod
    def save_waiting_json(info: Dict[str, dict], filename: str):
        with open(filename, 'w', encoding='utf-8') as f:
            json.dump(info, f, ensure_ascii=False, indent=2)
        logger.info(f"待上架信息已保存至 {filename}")


class PutawayProcessor:
    @staticmethod
    def get_rt_info(token: str, vtech_lot: str) -> Optional[dict]:
        global logger
        url = f"{GET_RT_INFO_URL}/{vtech_lot}/ON"
        headers = {
            "Accept": "application/json, text/plain, */*",
            "Accept-Encoding": "gzip, deflate",
            "Accept-Language": "zh-CN,zh;q=0.9",
            "Connection": "keep-alive",
            "Host": "10.97.245.205:92",
            "Referer": "http://10.97.245.205:92/wms/sendproduction/onshelf/index",
            "token": token,
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
        }
        for attempt in range(1, MAX_RETRY + 1):
            try:
                resp = requests.get(url, headers=headers, timeout=10)
                resp.raise_for_status()
                data = resp.json()
                if isinstance(data, list) and len(data) > 0:
                    return data[0]
                else:
                    logger.warning(f"getRTInfo 返回非预期格式: {data}")
                    return None
            except Exception as e:
                logger.warning(f"getRTInfo 第{attempt}次请求异常: {e}")
            if attempt < MAX_RETRY:
                time.sleep(RETRY_DELAY)
        return None

    @staticmethod
    def do_putaway(token: str, batch: str, bincode: str, qty: float, account: str) -> Tuple[bool, str]:
        """
        执行上架，返回 (是否成功, 响应消息)
        """
        global logger
        json_param = {
            "p_bincode": bincode,
            "p_batch": batch,
            "p_totalqty": qty,
            "p_inqty": 0,
            "p_qty": qty,
            "account": account
        }
        json_str = json.dumps(json_param, ensure_ascii=False)
        url = f"{PUTAWAY_URL}?spname=do_wms_batch_instore_sp&json={json_str}"
        headers = {
            "Accept": "application/json, text/plain, */*",
            "Accept-Encoding": "gzip, deflate",
            "Accept-Language": "zh-CN,zh;q=0.9",
            "Connection": "keep-alive",
            "Host": "10.97.245.205:92",
            "Origin": "http://10.97.245.205:92",
            "Referer": "http://10.97.245.205:92/wms/sendproduction/onshelf/index",
            "token": token,
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
        }
        for attempt in range(1, MAX_RETRY + 1):
            try:
                resp = requests.put(url, headers=headers, timeout=30)
                resp.raise_for_status()
                result = resp.json()
                if "Table" in result and len(result["Table"]) > 0:
                    row = result["Table"][0]
                    if row.get("o_res") == "OK":
                        return True, "成功"
                    else:
                        return False, f"返回错误: {row}"
                else:
                    return False, f"响应格式异常: {result}"
            except Exception as e:
                logger.warning(f"上架第{attempt}次请求异常: {e}")
            if attempt < MAX_RETRY:
                time.sleep(RETRY_DELAY)
        return False, "重试失败，网络异常"

    @staticmethod
    def process_putaway(
        waiting_info: Dict[str, dict],
        smart_stats: Dict[str, float],
        account: str,
        forced_location: str
    ) -> List[dict]:
        """
        执行所有上架，并返回每个 RT 的结果明细列表
        """
        global logger
        token = MESAuth.login(MES_ACCOUNT, MES_PASSWORD)
        if not token:
            logger.error("登录失败，无法上架")
            return []

        logger.info(f"\n待上架 RT 数量: {len(waiting_info)}")
        logger.info(f"强制库位: {forced_location}（所有上架均使用此库位，忽略报表中的 location）")

        results = []   # 每个元素是一个 dict，包含 RT、待上架数量、智能料架数量、实际上架数量、库位、状态、消息

        for vtech, wait_data in waiting_info.items():
            wait_qty = wait_data["total_qty"]
            shelf_qty = smart_stats.get(vtech, 0.0)
            put_qty = min(wait_qty, shelf_qty) if shelf_qty > 0 else 0

            logger.info(f"\n处理 RT: {vtech}")
            logger.info(f"  待上架数量: {wait_qty}, 智能货架总数量: {shelf_qty}, 实际上架: {put_qty}")

            result_item = {
                "vtech_lot": vtech,
                "wait_qty": wait_qty,
                "shelf_qty": shelf_qty,
                "put_qty": put_qty,
                "location": forced_location,
                "status": "",
                "message": ""
            }

            if put_qty <= 0:
                logger.info("  上架数量为 0，跳过")
                result_item["status"] = "跳过"
                result_item["message"] = "上架数量为0"
                results.append(result_item)
                continue

            # 可选：获取 RT 详情（仅用于日志）
            rt_info = PutawayProcessor.get_rt_info(token, vtech)
            if rt_info:
                logger.info(f"  RT Info: part_no={rt_info.get('part_no')}, unit_qty={rt_info.get('unit_qty')}")
            else:
                logger.warning("  警告：获取 RT 详情失败，继续上架")

            # 执行上架
            success, msg = PutawayProcessor.do_putaway(token, vtech, forced_location, put_qty, account)
            if success:
                logger.info(f"  上架成功: batch={vtech}, qty={put_qty}")
                result_item["status"] = "成功"
                result_item["message"] = msg
            else:
                logger.error(f"  上架失败: {msg}")
                result_item["status"] = "失败"
                result_item["message"] = msg
            results.append(result_item)

        # 汇总统计
        success_count = sum(1 for r in results if r["status"] == "成功")
        fail_count = sum(1 for r in results if r["status"] == "失败")
        skip_count = sum(1 for r in results if r["status"] == "跳过")
        logger.info(f"\n上架完成：成功 {success_count}，失败 {fail_count}，跳过 {skip_count}")
        return results

    @staticmethod
    def save_result_to_excel(results: List[dict], forced_location: str):
        """生成 Excel 结果报表"""
        global logger
        if not results:
            logger.warning("没有结果数据，不生成报表")
            return

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"putaway_result_{timestamp}.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "上架结果"

        # 表头
        headers = ["RT", "待上架数量", "智能料架总数量", "实际上架数量", "上架库位", "状态", "消息"]
        ws.append(headers)

        # 样式
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

        for r in results:
            row = [
                r["vtech_lot"],
                r["wait_qty"],
                r["shelf_qty"],
                r["put_qty"],
                r["location"],
                r["status"],
                r["message"]
            ]
            ws.append(row)

            # 根据状态着色
            row_idx = ws.max_row
            status_cell = ws.cell(row=row_idx, column=6)
            if r["status"] == "成功":
                status_cell.fill = green_fill
            elif r["status"] == "失败":
                status_cell.fill = red_fill
            elif r["status"] == "跳过":
                status_cell.fill = yellow_fill

        # 调整列宽
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws.column_dimensions[col_letter].width = adjusted_width

        wb.save(filename)
        logger.info(f"上架结果报表已保存至 {filename}")

        # 同时保存为 JSON 备份
        json_filename = f"putaway_result_{timestamp}.json"
        with open(json_filename, 'w', encoding='utf-8') as f:
            json.dump(results, f, ensure_ascii=False, indent=2)
        logger.info(f"结果 JSON 备份保存至 {json_filename}")


# ======================== 交互式主程序 ================================

def main():
    global logger
    logger = setup_logger()

    print("=" * 50)
    print("        MES 自动上架工具 - 交互式菜单")
    print("=" * 50)
    print("  1. 下载智能货架数据 (步骤1)")
    print("  2. 下载 MES 待上架报表 (步骤2)")
    print("  3. 执行上架 (步骤3) - 强制使用 SA00 库位")
    print("  4. 全流程 (步骤1+2+3)")
    print("  0. 退出")
    print("=" * 50)

    choice = input("请输入数字 (0-4): ").strip()
    if choice not in ['0','1','2','3','4']:
        logger.error("输入无效，请重新运行。")
        input("按 Enter 键退出...")
        sys.exit(1)

    if choice == '0':
        logger.info("退出程序")
        sys.exit(0)

    forced_loc = FORCED_LOCATION
    logger.info(f"注意：所有上架将强制使用库位 '{forced_loc}'，不受报表影响。")

    # ---------- 步骤1 ----------
    if choice in ['1', '4']:
        logger.info("\n===== 步骤1：下载智能货架 =====")
        print("请确保已连接到智能货架网络 (192.168.9.100)")
        input("按 Enter 键继续...")
        smart_stats = SmartShelfDownloader.download_and_stats()
        SmartShelfDownloader.save_stats_to_json(smart_stats, SMART_SHELF_JSON)
        logger.info("步骤1完成\n")
        if choice == '1':
            input("按 Enter 键退出...")
            sys.exit(0)

    # ---------- 步骤2 ----------
    if choice in ['2', '4']:
        logger.info("\n===== 步骤2：下载 MES 待上架报表 =====")
        print("请确保已连接到 MES 网络 (10.97.245.205)")
        input("按 Enter 键继续...")
        token = MESAuth.login(MES_ACCOUNT, MES_PASSWORD)
        if not token:
            logger.error("登录失败，退出")
            input("按 Enter 键退出...")
            sys.exit(1)
        rows = WaitingReportDownloader.download_all(token)
        if rows:
            WaitingReportDownloader.save_to_excel(rows, WAITING_REPORT_EXCEL)
            waiting_info = WaitingReportDownloader.extract_waiting_info(rows)
            WaitingReportDownloader.save_waiting_json(waiting_info, WAITING_REPORT_JSON)
        else:
            logger.warning("未获取到任何待上架数据")
        logger.info("步骤2完成\n")
        if choice == '2':
            input("按 Enter 键退出...")
            sys.exit(0)

    # ---------- 步骤3 ----------
    if choice in ['3', '4']:
        logger.info("\n===== 步骤3：执行上架 =====")
        print("请确保已连接到 MES 网络 (10.97.245.205)")
        input("按 Enter 键继续...")

        if not os.path.exists(SMART_SHELF_JSON):
            logger.error(f"{SMART_SHELF_JSON} 不存在，请先执行步骤1")
            input("按 Enter 键退出...")
            sys.exit(1)
        if not os.path.exists(WAITING_REPORT_JSON):
            logger.error(f"{WAITING_REPORT_JSON} 不存在，请先执行步骤2")
            input("按 Enter 键退出...")
            sys.exit(1)

        with open(SMART_SHELF_JSON, 'r', encoding='utf-8') as f:
            smart_stats = json.load(f)
        with open(WAITING_REPORT_JSON, 'r', encoding='utf-8') as f:
            waiting_info = json.load(f)

        results = PutawayProcessor.process_putaway(waiting_info, smart_stats, MES_ACCOUNT, forced_loc)
        if results:
            PutawayProcessor.save_result_to_excel(results, forced_loc)
        else:
            logger.warning("没有产生任何上架结果")

        logger.info("步骤3完成\n")
        input("按 Enter 键退出...")


if __name__ == "__main__":
    main()