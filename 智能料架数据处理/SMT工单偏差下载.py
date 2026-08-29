import requests
from openpyxl import Workbook
from openpyxl.styles import PatternFill
import time

def main():
    print("=== 订单欠料详情下载工具 ===")
    order_id = input("请输入订单号（如 PASCAL_8466to9265）：").strip()
    if not order_id:
        print("订单号不能为空，退出")
        return

    url = "http://192.168.9.100:8081/order/orderDetail"

    headers = {
        "Accept": "application/json, text/plain, */*",
        "Accept-Encoding": "gzip, deflate",
        "Accept-Language": "zh-CN,zh;q=0.9",
        "Connection": "keep-alive",
        "Content-Type": "application/json;charset=UTF-8",
        "Host": "192.168.9.100:8081",
        "Origin": "http://192.168.9.100:8081",
        "Referer": "http://192.168.9.100:8081/",
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/151.0.0.0 Safari/537.36"
    }

    payload = {
        "order_id": order_id,
        "user_name": "NAF684",
        "user_group_power": "1",
        "user_language": "zh"
    }

    # 重试机制
    max_retries = 3
    retry_delay = 5
    resp = None
    for attempt in range(1, max_retries + 1):
        try:
            print(f"正在请求（尝试 {attempt}/{max_retries}）...")
            resp = requests.post(url, json=payload, headers=headers, timeout=30)
            break
        except requests.exceptions.Timeout:
            if attempt == max_retries:
                print(f"请求超时，已重试 {max_retries} 次，请检查网络或服务状态。")
                return
            print(f"超时，{retry_delay}秒后重试...")
            time.sleep(retry_delay)
        except Exception as e:
            print(f"请求异常：{e}")
            return

    if resp is None:
        print("未获得响应")
        return

    if resp.status_code != 200:
        print(f"请求失败（状态码 {resp.status_code}）")
        return

    data = resp.json()
    if data.get("result") != 1:
        print(f"接口返回 result 非 1，返回内容：{data}")
        return

    items = data.get("data", [])
    if not items:
        print("无数据")
        return

    print(f"成功获取 {len(items)} 条记录")

    # 准备写入 Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "订单欠料"

    # 定义表头（按需排序）
    headers = [
        "id", "order_num", "feeder_name", "lot_code", "master_pn",
        "slave_pn", "need_num", "issue_quantity", "欠料",
        "mate_state", "issue_state", "send_num", "send_mate_list"
    ]

    # 写入表头
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    # 红色填充样式
    red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

    # 写入数据行
    for row_idx, item in enumerate(items, start=2):
        # 解析字段
        row_data = {
            "id": item.get("id"),
            "order_num": item.get("order_num"),
            "feeder_name": item.get("feeder_name"),
            "lot_code": item.get("lot_code"),
            "master_pn": item.get("master_pn"),
            "slave_pn": str(item.get("slave_pn", [])),
            "need_num": item.get("need_num"),
            "issue_quantity": item.get("issue_quantity"),
            "mate_state": item.get("mate_state"),
            "issue_state": item.get("issue_state"),
            "send_num": item.get("send_num"),
            "send_mate_list": str(item.get("send_mate_list", []))
        }
        # 计算欠料
        need = row_data["need_num"] if row_data["need_num"] is not None else 0
        issued = row_data["issue_quantity"] if row_data["issue_quantity"] is not None else 0
        shortage = issued - need   # 负数表示欠料（已发少于需求）
        row_data["欠料"] = shortage

        # 写入所有列
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row_data.get(header))
            # 如果当前列是"欠料"且值为负数，标红
            if header == "欠料" and shortage < 0:
                cell.fill = red_fill

    # 保存文件
    out_file = f"订单欠料_{order_id}.xlsx"
    wb.save(out_file)
    print(f"数据已保存至 {out_file}，共 {len(items)} 条记录")

if __name__ == "__main__":
    main()PASCAL_8466to9265