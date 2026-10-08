# -*- coding: utf-8 -*-
"""
Cummins / Crane BOM & Actual Issue Tool
=========================================

双语界面 / Bilingual GUI (English + 中文)
拖拽 Excel / Drag & drop Excel
读取 BOM / Read BOM
查询生产在库 / Fetch production stock
生成生产BOM / 生成缺料清单 / 回写实际发料
Generate production BOM / actual issue / write back actual issued

生产在库仓库 ID 从独立配置文件读取（可界面修改并保存）。
Production warehouse ID is read from a standalone config file
(editable & saved from GUI). Semantics: source of PRODUCTION STOCK.

模式 / Modes:
  (1) Generate Production BOM   : 全部生产BOM，不扣在库 / full BOM, no stock deduction
  (2) Generate Actual Issue     : 需求 − 生产在库 / demand − production stock
  (3) Write Back Actual Issued  : 生产在库 + 订单API 写入 E 列，F 列偏差
                                  production stock + order API into E, deviation into F

  (1) 和 (2) 二选一 / (1) and (2) are mutually exclusive (radio buttons).

UI style: white + blue, scrollable body
"""

import os
import re
import json
import time
import threading
import subprocess
import traceback
from datetime import datetime
from typing import Tuple, List

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from tkinterdnd2 import TkinterDnD, DND_FILES

import requests
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill

# =========================================================
# API config
# =========================================================
API_URL = "http://192.168.9.100:8081/mate/searchByInfo"
PAGE_SIZE = 10000
USER_NAME = "NAF684"
USER_GROUP_POWER = "1"
USER_LANGUAGE = "zh"
MAX_RETRY = 3
RETRY_DELAY = 2

ORDER_API_URL = "http://192.168.9.100:8081/order/orderDetail"

ORDER_HEADERS = {
    "Accept": "application/json, text/plain, */*",
    "Accept-Encoding": "gzip, deflate",
    "Accept-Language": "zh-CN,zh;q=0.9",
    "Connection": "keep-alive",
    "Content-Type": "application/json;charset=UTF-8",
    "Host": "192.168.9.100:8081",
    "Origin": "http://192.168.9.100:8081",
    "Referer": "http://192.168.9.100:8081/",
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/151.0.0.0 Safari/537.36"
    ),
}

# =========================================================
# WiFi config (reused from Batch Outbound Tool)
# =========================================================
WIFI_PROFILE_NAME = "AOI"
WIFI_PASSWORD = "123456789"
WIFI_PROFILE_XML = f"""<?xml version="1.0"?>
<WLANProfile xmlns="http://www.microsoft.com/networking/WLAN/profile/v1">
    <name>{WIFI_PROFILE_NAME}</name>
    <SSIDConfig>
        <SSID>
            <name>{WIFI_PROFILE_NAME}</name>
        </SSID>
    </SSIDConfig>
    <connectionType>ESS</connectionType>
    <connectionMode>auto</connectionMode>
    <MSM>
        <security>
            <authEncryption>
                <authentication>WPA2PSK</authentication>
                <encryption>AES</encryption>
                <useOneX>false</useOneX>
            </authEncryption>
            <sharedKey>
                <keyType>passPhrase</keyType>
                <protected>false</protected>
                <keyMaterial>{WIFI_PASSWORD}</keyMaterial>
            </sharedKey>
        </security>
    </MSM>
</WLANProfile>
"""


# =========================================================
# Config file (standalone)
# =========================================================
CONFIG_FILENAME = "bom_tool_config.json"
DEFAULT_PRODUCTION_WAREHOUSE_ID = 39


def config_path() -> str:
    try:
        base = os.path.dirname(os.path.abspath(__file__))
    except Exception:
        base = os.getcwd()
    return os.path.join(base, CONFIG_FILENAME)


def load_config() -> dict:
    default = {"production_warehouse_id": DEFAULT_PRODUCTION_WAREHOUSE_ID}
    path = config_path()
    if not os.path.isfile(path):
        return default
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        try:
            data["production_warehouse_id"] = int(
                data.get("production_warehouse_id",
                         DEFAULT_PRODUCTION_WAREHOUSE_ID))
        except (TypeError, ValueError):
            data["production_warehouse_id"] = DEFAULT_PRODUCTION_WAREHOUSE_ID
        return data
    except Exception:
        return default


def save_config(cfg: dict) -> bool:
    try:
        with open(config_path(), "w", encoding="utf-8") as f:
            json.dump(cfg, f, ensure_ascii=False, indent=2)
        return True
    except Exception:
        return False


def parse_single_warehouse_id(text) -> int:
    if text is None:
        raise ValueError("Warehouse ID is empty / 仓库ID为空")
    s = str(text).strip()
    if not s:
        raise ValueError("Warehouse ID is empty / 仓库ID为空")
    tokens = re.split(r"[\s,;，、]+", s)
    for t in tokens:
        if not t:
            continue
        try:
            return int(t)
        except ValueError:
            raise ValueError(f"Invalid warehouse ID / 无效仓库ID: {t}")
    raise ValueError("No valid warehouse ID / 未找到有效仓库ID")


# =========================================================
# Color scheme (white + blue)
# =========================================================
COLOR_BG        = "#FFFFFF"
COLOR_BG_ALT    = "#F5F9FF"
COLOR_PRIMARY   = "#1976D2"
COLOR_PRIMARY_D = "#0D47A1"
COLOR_ACCENT    = "#E3F2FD"
COLOR_BORDER    = "#BBDEFB"
COLOR_TEXT      = "#1A1A1A"
COLOR_MUTED     = "#607D8B"
COLOR_SUCCESS   = "#2E7D32"
COLOR_WARN      = "#EF6C00"
COLOR_ERROR     = "#C62828"

FONT_BASE  = ("Segoe UI", 10)
FONT_BOLD  = ("Segoe UI", 10, "bold")
FONT_TITLE = ("Segoe UI", 15, "bold")
FONT_SUB   = ("Segoe UI", 9)
FONT_MONO  = ("Consolas", 9)


# =========================================================
# Helpers
# =========================================================
def safe_float(v):
    try:
        if v is None or v == "":
            return 0.0
        return float(v)
    except Exception:
        return 0.0


def safe_int(v):
    try:
        if v is None or v == "":
            return 0
        return int(float(v))
    except Exception:
        return 0


# =========================================================
# WiFi Manager (reused from Batch Outbound Tool)
# =========================================================
class WifiManager:
    @staticmethod
    def _run(args, timeout=15) -> str:
        try:
            out = subprocess.run(args, capture_output=True, timeout=timeout)
            return (out.stdout or b"").decode("gbk", errors="ignore")
        except Exception:
            return ""

    @staticmethod
    def get_current_ssid() -> str:
        out = WifiManager._run(["netsh", "wlan", "show", "interfaces"])
        for line in out.splitlines():
            if "SSID" in line and "BSSID" not in line:
                parts = line.split(":", 1)
                if len(parts) == 2:
                    return parts[1].strip()
        return ""

    @staticmethod
    def _get_state_and_ssid() -> Tuple[str, str]:
        out = WifiManager._run(["netsh", "wlan", "show", "interfaces"])
        state = ""
        ssid = ""
        for line in out.splitlines():
            if "State" in line or "状态" in line:
                if ":" in line:
                    state = line.split(":", 1)[1].strip().lower()
            if "SSID" in line and "BSSID" not in line:
                if ":" in line:
                    ssid = line.split(":", 1)[1].strip()
        return state, ssid

    @staticmethod
    def list_profiles() -> List[str]:
        out = WifiManager._run(["netsh", "wlan", "show", "profiles"])
        profiles = []
        for line in out.splitlines():
            if ":" in line and ("All User Profile" in line
                                or "所有用户配置文件" in line
                                or "User Profile" in line
                                or "用户配置文件" in line):
                profiles.append(line.split(":", 1)[1].strip())
        return profiles

    @staticmethod
    def ensure_profile() -> bool:
        profiles = WifiManager.list_profiles()
        if WIFI_PROFILE_NAME in profiles:
            return True
        try:
            from pathlib import Path
            xml_path = Path(__file__).resolve().parent / "_wifi_profile.xml"
            xml_path.write_text(WIFI_PROFILE_XML, encoding="utf-8")
            WifiManager._run(
                ["netsh", "wlan", "add", "profile",
                 f"filename={xml_path}", "user=all"],
                timeout=15
            )
            try:
                xml_path.unlink()
            except Exception:
                pass
            return WIFI_PROFILE_NAME in WifiManager.list_profiles()
        except Exception:
            return False

    @staticmethod
    def disconnect():
        WifiManager._run(["netsh", "wlan", "disconnect"], timeout=10)
        time.sleep(1.5)

    @staticmethod
    def connect(profile_name: str, timeout: int = 30) -> bool:
        state, ssid = WifiManager._get_state_and_ssid()
        if ssid == profile_name and state == "connected":
            return True

        WifiManager.disconnect()

        try:
            subprocess.run(
                ["netsh", "wlan", "connect", f"name={profile_name}"],
                check=False, timeout=15
            )
        except Exception:
            return False

        t0 = time.time()
        last_retry = t0
        while time.time() - t0 < timeout:
            state, ssid = WifiManager._get_state_and_ssid()
            if ssid == profile_name and state == "connected":
                time.sleep(1.5)
                state2, ssid2 = WifiManager._get_state_and_ssid()
                if ssid2 == profile_name and state2 == "connected":
                    return True
            if time.time() - last_retry > 15:
                last_retry = time.time()
                try:
                    subprocess.run(
                        ["netsh", "wlan", "connect", f"name={profile_name}"],
                        check=False, timeout=15
                    )
                except Exception:
                    pass
            time.sleep(0.5)
        return False


# =========================================================
# Custom Order ID Dialog (white + blue style)
# =========================================================
class OrderIdDialog(tk.Toplevel):
    """白蓝风格 Order ID 输入弹窗 / Blue-white style Order ID dialog."""

    def __init__(self, parent, default_value=""):
        super().__init__(parent)
        self.title("Order ID  订单号")
        self.configure(bg=COLOR_BG)
        self.resizable(False, False)
        self.result = None

        self.transient(parent)
        self.grab_set()

        body = ttk.Frame(self, style="TFrame", padding=(20, 18, 20, 14))
        body.pack(fill=tk.BOTH, expand=True)

        ttk.Label(
            body,
            text="请输入订单号 / Enter order ID:",
            background=COLOR_BG, foreground=COLOR_TEXT, font=FONT_BOLD
        ).pack(anchor="w", pady=(0, 4))

        ttk.Label(
            body,
            text="例如 / e.g.  PASCAL_8466to9265",
            style="Muted.TLabel"
        ).pack(anchor="w", pady=(0, 10))

        self.entry_var = tk.StringVar(value=default_value)
        entry = ttk.Entry(body, textvariable=self.entry_var, width=48)
        entry.pack(fill=tk.X, pady=(0, 14))
        entry.focus_set()
        entry.select_range(0, tk.END)

        btn_row = ttk.Frame(body, style="TFrame")
        btn_row.pack(fill=tk.X)

        ttk.Button(btn_row, text="Cancel  取消",
                   style="Secondary.TButton",
                   command=self._on_cancel).pack(side=tk.RIGHT, padx=(8, 0))
        ttk.Button(btn_row, text="OK  确定",
                   style="Primary.TButton",
                   command=self._on_ok).pack(side=tk.RIGHT)

        self.bind("<Return>", lambda e: self._on_ok())
        self.bind("<Escape>", lambda e: self._on_cancel())

        self.update_idletasks()
        px = parent.winfo_rootx()
        py = parent.winfo_rooty()
        pw = parent.winfo_width()
        ph = parent.winfo_height()
        w = self.winfo_width()
        h = self.winfo_height()
        x = px + (pw - w) // 2
        y = py + (ph - h) // 2
        self.geometry(f"+{max(0, x)}+{max(0, y)}")

        self.wait_window(self)

    def _on_ok(self):
        v = self.entry_var.get().strip()
        if not v:
            messagebox.showwarning(
                "Empty  为空",
                "Order ID cannot be empty.\n订单号不能为空。",
                parent=self)
            return
        self.result = v
        self.destroy()

    def _on_cancel(self):
        self.result = None
        self.destroy()


# =========================================================
# Warehouse API
# =========================================================
def build_payload(start, rows, warehouse_name_id=None):
    payload = {
        "part_num": "",
        "warehouse_category_id": 0,
        "warehouse_name_id": 0,
        "save_id": "",
        "lot_code": "",
        "mfg_date_start": "",
        "supplier_name": "",
        "mfg_date_end": "",
        "start": start,
        "rows": rows,
        "user_name": USER_NAME,
        "user_group_power": USER_GROUP_POWER,
        "user_language": USER_LANGUAGE,
        "shelf_id": "",
        "position_info": "",
        "start_date": "",
        "end_date": "",
    }
    if warehouse_name_id is not None:
        payload["warehouse_name_id"] = warehouse_name_id
    return payload


def fetch_page(start, warehouse_name_id=None, logger=None):
    payload = build_payload(start, PAGE_SIZE, warehouse_name_id=warehouse_name_id)
    for attempt in range(1, MAX_RETRY + 1):
        try:
            resp = requests.post(API_URL, json=payload, timeout=30)
            resp.raise_for_status()
            data = resp.json()
            if data.get("result") == 1:
                return data
            if logger:
                logger(f"[WARN] attempt {attempt}: result != 1: {data}")
        except Exception as e:
            if logger:
                logger(f"[WARN] attempt {attempt}: request error: {e}")
        if attempt < MAX_RETRY:
            time.sleep(RETRY_DELAY)
    return None


def fetch_all_warehouse_by_id(warehouse_name_id, log_func=None):
    def log(msg):
        if log_func:
            log_func(msg)
        else:
            print(msg)

    all_raw = []
    start = 0
    total = None
    consecutive_empty = 0

    while True:
        data = fetch_page(start, warehouse_name_id=warehouse_name_id, logger=log)
        if data is None:
            log(f"[ERROR] start={start} failed after retries, stop.")
            break

        mate_list = data.get("mate_list", [])
        if total is None:
            total = data.get("rows", 0)
            log(f"API total rows (warehouse_name_id={warehouse_name_id}): {total}")

        if not mate_list:
            consecutive_empty += 1
            log(f"[INFO] start={start} empty page ({consecutive_empty}/3)")
            if consecutive_empty >= 3:
                break
            start += PAGE_SIZE
            time.sleep(0.5)
            continue

        consecutive_empty = 0
        all_raw.extend(mate_list)
        log(f"Fetched {len(all_raw)} records ... start={start}")

        if total > 0 and len(all_raw) > total * 1.5:
            log("[WARN] fetched more than 1.5x declared total, stop.")
            break

        start += PAGE_SIZE
        time.sleep(0.1)

    return all_raw


def build_stock_map(warehouse_records):
    stock = {}
    for rec in warehouse_records:
        part = str(rec.get("part_num", "") or "").strip()
        if not part:
            continue
        qty = safe_float(rec.get("quantity", 0))
        stock[part] = stock.get(part, 0.0) + qty
    return stock


# =========================================================
# Order Detail API
# =========================================================
def fetch_order_detail(order_id):
    payload = {
        "order_id": order_id,
        "user_name": USER_NAME,
        "user_group_power": USER_GROUP_POWER,
        "user_language": USER_LANGUAGE,
    }

    max_retries = 3
    retry_delay = 5
    resp = None

    for attempt in range(1, max_retries + 1):
        try:
            print(f"Order detail request (attempt {attempt}/{max_retries}) ...")
            resp = requests.post(ORDER_API_URL, json=payload,
                                 headers=ORDER_HEADERS, timeout=30)
            break
        except requests.exceptions.Timeout:
            if attempt == max_retries:
                raise TimeoutError(
                    f"Order detail request timeout after {max_retries} retries.")
            print(f"Timeout, retry in {retry_delay}s ...")
            time.sleep(retry_delay)
        except Exception as e:
            raise RuntimeError(f"Order detail request error: {e}")

    if resp is None:
        raise RuntimeError("No response from order detail API.")

    if resp.status_code != 200:
        raise RuntimeError(f"Order detail API HTTP {resp.status_code}")

    data = resp.json()
    if data.get("result") != 1:
        raise RuntimeError(f"Order detail API result != 1: {data}")

    items = data.get("data", [])
    return items


def build_issued_map_from_order(items):
    issued_map = {}
    for it in items:
        part = str(it.get("master_pn", "") or "").strip()
        if not part:
            continue
        qty = safe_float(it.get("issue_quantity", 0))
        issued_map[part] = issued_map.get(part, 0.0) + qty
    return issued_map


# =========================================================
# Read BOM Excel
# =========================================================
def read_bom_excel(file_path):
    wb = load_workbook(file_path, data_only=True)
    ws = wb.active

    bom_title = ""
    for row in ws.iter_rows(min_row=1, max_row=10, max_col=8, values_only=True):
        for cell in row:
            if cell and isinstance(cell, str) and "Bill of Materials" in cell:
                bom_title = cell.strip()
                break
        if bom_title:
            break

    header_row_idx = None
    for i, row in enumerate(
            ws.iter_rows(min_row=1, max_row=40, max_col=8, values_only=True),
            start=1):
        row_text = " ".join([str(c) for c in row if c is not None])
        if (("物料编号" in row_text or "Part No" in row_text)
                and ("总需求数" in row_text or "Total Qty" in row_text)):
            header_row_idx = i
            break

    if header_row_idx is None:
        raise ValueError(
            "Cannot find BOM item header row (物料编号 / 总需求数).\n"
            "找不到 BOM 物料表头行。")

    items = []
    no = 0
    for row in ws.iter_rows(min_row=header_row_idx + 1, max_row=ws.max_row,
                            max_col=8, values_only=True):
        part_no = row[1]
        if not part_no:
            continue
        part_no = str(part_no).strip()
        if not part_no or part_no.lower() in ("none", "null"):
            continue
        if "物料编号" in part_no or "Part No" in part_no:
            continue

        location = row[2] if len(row) > 2 and row[2] is not None else ""
        total_qty = row[3] if len(row) > 3 else 0

        no += 1
        items.append({
            "no": no,
            "part_no": part_no,
            "location": str(location).strip(),
            "total_qty": safe_float(total_qty),
        })

    if not items:
        raise ValueError("No BOM items found. / 未找到 BOM 物料。")

    return bom_title, items, header_row_idx, ws, wb


# =========================================================
# Mode 1: Generate BOM table
# =========================================================
def generate_bom_table(items):
    rows = []
    for idx, it in enumerate(items, start=1):
        rows.append([
            idx,
            it["part_no"],
            "",
            "",
            it["total_qty"],
            it["location"],
        ])
    return rows


# =========================================================
# Mode 2: Generate actual issue table
# =========================================================
def generate_actual_issue_table(bom_items, production_stock_map, log_func=None):
    def log(msg):
        if log_func:
            log_func(msg)
        else:
            print(msg)

    log("========== Compare Log: BOM vs Production Stock ==========")
    rows = []
    idx = 0
    skipped_no_stock = []
    skipped_enough = []
    need_issue = []

    for it in bom_items:
        part = it["part_no"]
        demand = safe_float(it["total_qty"])
        stock = production_stock_map.get(part, None)

        if stock is None:
            log(f"[NO STOCK RECORD] Part={part} | Demand={demand} | "
                f"ProdStock=NOT FOUND -> treat as 0")
            stock = 0.0
            skipped_no_stock.append(part)
        else:
            log(f"[OK] Part={part} | Demand={demand} | ProdStock={stock}")

        actual_issue = demand - stock

        if actual_issue <= 0:
            log(f"    -> SKIP (stock enough) | Actual Issue = {actual_issue}")
            skipped_enough.append(part)
            continue

        idx += 1
        rows.append([idx, part, "", "", actual_issue, ""])
        need_issue.append(part)
        log(f"    -> NEED ISSUE | Actual Issue = {actual_issue}")

    log("========== Compare Summary ==========")
    log(f"Total BOM items        : {len(bom_items)}")
    log(f"Need issue items       : {len(need_issue)}")
    log(f"Skip (stock enough)    : {len(skipped_enough)}")
    log(f"No production stock    : {len(skipped_no_stock)}")
    if skipped_no_stock:
        log("Parts with NO production stock record:")
        for p in skipped_no_stock:
            log(f"    - {p}")
    if skipped_enough:
        log("Parts skipped because stock enough:")
        for p in skipped_enough:
            log(f"    - {p}")
    log("========== End Compare Log ==========")

    return rows


# =========================================================
# Excel export
# =========================================================
def export_bom_excel(rows, output_path, title=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"

    headers = ["序号", "物料编号", "物料名称", "规格型号", "单位用量", "位置号"]
    ws.append(headers)

    for r in rows:
        ws.append(r)

    widths = [8, 24, 18, 18, 12, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    wb.save(output_path)
    return output_path


def export_actual_issue_excel(rows, output_path, title=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Actual Issue"

    headers = ["序号", "物料编号", "物料名称", "规格型号", "实际发料", "位置号"]
    ws.append(headers)

    for r in rows:
        ws.append(r)

    widths = [8, 24, 18, 18, 14, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    wb.save(output_path)
    return output_path


# =========================================================
# Mode 3: Write back actual issued + deviation
# =========================================================
def write_back_actual_issued(src_path, issued_map, production_stock_map,
                             out_path, log_func=None):
    """
    回写 BOM Excel / Write back to BOM Excel:
      E 列（实际发料）逻辑：
        两个都有 -> E = 生产在库 + issue_quantity, 标浅蓝
        只有 issue_quantity -> E = issue_quantity, 标绿
        只有生产在库 -> E = 生产在库, 不标色
        都没有 -> 留空
      F 列（偏差）= E − D，F < 0 -> 只 F 单元格标淡红
      I 列及之后不写
    """
    def log(msg):
        if log_func:
            log_func(msg)
        else:
            print(msg)

    wb = load_workbook(src_path)
    ws = wb.active

    # ---- 找表头行 ----
    header_row_idx = None
    for i, row in enumerate(
            ws.iter_rows(min_row=1, max_row=40, max_col=8, values_only=True),
            start=1):
        row_text = " ".join([str(c) for c in row if c is not None])
        if (("物料编号" in row_text or "Part No" in row_text)
                and ("总需求数" in row_text or "Total Qty" in row_text)):
            header_row_idx = i
            break

    if header_row_idx is None:
        raise ValueError(
            "Cannot find BOM item header row for write-back.\n"
            "找不到 BOM 物料表头行。")

    # ---- 颜色 ----
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE",
                             fill_type="solid")   # 绿
    blue_fill  = PatternFill(start_color="DDEBF7", end_color="DDEBF7",
                             fill_type="solid")   # 浅蓝
    red_fill   = PatternFill(start_color="FFC7CE", end_color="FFC7CE",
                             fill_type="solid")   # 淡红

    log("========== Write Back Compare Log ==========")
    both_cnt = 0
    only_issue_cnt = 0
    only_stock_cnt = 0
    blank_cnt = 0
    negative_deviation = 0

    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        part_cell = ws.cell(row=row_idx, column=2)   # B
        part = part_cell.value
        if not part:
            continue
        part = str(part).strip()
        if not part or part.lower() in ("none", "null"):
            continue
        if "物料编号" in part or "Part No" in part:
            continue

        e_cell = ws.cell(row=row_idx, column=5)   # E
        d_cell = ws.cell(row=row_idx, column=4)   # D
        f_cell = ws.cell(row=row_idx, column=6)   # F

        d_val = safe_float(d_cell.value)

        stock_qty = production_stock_map.get(part, 0.0)
        issue_qty = issued_map.get(part, 0.0)

        has_stock = stock_qty > 0
        has_issue = issue_qty > 0

        # ---- E 列：根据情况填写 ----
        if has_stock and has_issue:
            e_val = stock_qty + issue_qty
            e_cell.value = e_val
            e_cell.fill = blue_fill
            both_cnt += 1
            log(f"[E] Part={part} | Stock={stock_qty} + Issue={issue_qty} "
                f"= {e_val} -> E{row_idx} (BLUE)")
        elif has_issue:
            e_val = issue_qty
            e_cell.value = e_val
            e_cell.fill = green_fill
            only_issue_cnt += 1
            log(f"[E] Part={part} | Issue={issue_qty} -> E{row_idx} (GREEN)")
        elif has_stock:
            e_val = stock_qty
            e_cell.value = e_val
            only_stock_cnt += 1
            log(f"[E] Part={part} | Stock={stock_qty} -> E{row_idx} (no fill)")
        else:
            e_cell.value = None
            e_val = 0.0
            blank_cnt += 1
            log(f"[E] Part={part} | N/A -> E{row_idx} (blank)")

        # ---- F 列：偏差 = E − D ----
        deviation = e_val - d_val
        f_cell.value = deviation
        if deviation < 0:
            f_cell.fill = red_fill
            negative_deviation += 1
            log(f"[F] Part={part} | Deviation={deviation} -> F{row_idx} (RED)")
        else:
            log(f"[F] Part={part} | Deviation={deviation} -> F{row_idx}")

    log("========== Write Back Summary ==========")
    log(f"E both (stock + issue)  : {both_cnt} (blue)")
    log(f"E only issue            : {only_issue_cnt} (green)")
    log(f"E only stock            : {only_stock_cnt} (no fill)")
    log(f"E blank                 : {blank_cnt}")
    log(f"Negative deviation rows : {negative_deviation}")
    log("========== End Write Back Log ==========")

    wb.save(out_path)
    return both_cnt + only_issue_cnt + only_stock_cnt


# =========================================================
# GUI
# =========================================================
class App:
    def __init__(self, root):
        self.root = root
        self.root.title(
            "Cummins / Crane BOM & Actual Issue Tool  生产BOM / 缺料 / 回写工具")
        self.root.configure(bg=COLOR_BG)
        self._setup_styles()
        self._center_window(1100, 820)
        self.root.minsize(900, 600)

        self.config_data = load_config()
        self.production_warehouse_id = int(self.config_data.get(
            "production_warehouse_id", DEFAULT_PRODUCTION_WAREHOUSE_ID))

        self.file_path = tk.StringVar()
        self.status_var = tk.StringVar(
            value="Drag & drop an Excel file here, or click Browse.  "
                  "拖拽 Excel 文件到此处，或点击 Browse。")

        self.bom_items = []
        self.bom_title = ""
        self.base_name = ""
        self.production_stock = {}

        self._original_wifi = ""
        self._busy = False

        self._build_ui()
        self._refresh_wifi_label()

    # ---------------- Styles ----------------
    def _setup_styles(self):
        s = ttk.Style()
        try:
            s.theme_use("clam")
        except Exception:
            pass

        s.configure(".", background=COLOR_BG, foreground=COLOR_TEXT, font=FONT_BASE)
        s.configure("TFrame", background=COLOR_BG)
        s.configure("TLabel", background=COLOR_BG, foreground=COLOR_TEXT, font=FONT_BASE)
        s.configure("Muted.TLabel", background=COLOR_BG,
                    foreground=COLOR_MUTED, font=FONT_SUB)
        s.configure("Status.TLabel", background=COLOR_BG_ALT,
                    foreground=COLOR_PRIMARY_D, font=FONT_BOLD)
        s.configure("Warn.TLabel", background=COLOR_BG_ALT,
                    foreground=COLOR_WARN, font=FONT_BOLD)

        s.configure("Card.TLabelframe", background=COLOR_BG_ALT,
                    bordercolor=COLOR_BORDER, relief="solid", borderwidth=1)
        s.configure("Card.TLabelframe.Label", background=COLOR_BG_ALT,
                    foreground=COLOR_PRIMARY_D, font=FONT_BOLD)

        s.configure("TEntry", fieldbackground=COLOR_BG, foreground=COLOR_TEXT,
                    bordercolor=COLOR_BORDER, insertcolor=COLOR_PRIMARY)

        s.configure("TRadiobutton", background=COLOR_BG_ALT,
                    foreground=COLOR_TEXT, font=FONT_BOLD)

        s.configure("Primary.TButton",
                    background=COLOR_PRIMARY, foreground="white",
                    font=FONT_BOLD, padding=(12, 8),
                    borderwidth=0, focusthickness=0)
        s.map("Primary.TButton",
              background=[("active", COLOR_PRIMARY_D), ("pressed", COLOR_PRIMARY_D)],
              foreground=[("active", "white")])

        s.configure("Secondary.TButton",
                    background=COLOR_ACCENT, foreground=COLOR_PRIMARY_D,
                    font=FONT_BASE, padding=(10, 7),
                    borderwidth=0, focusthickness=0)
        s.map("Secondary.TButton",
              background=[("active", COLOR_BORDER), ("pressed", COLOR_BORDER)],
              foreground=[("active", COLOR_PRIMARY_D)])

        # 滚动条样式 / Slim scrollbar
        s.configure("Slim.Vertical.TScrollbar",
                    background=COLOR_ACCENT, troughcolor=COLOR_BG,
                    bordercolor=COLOR_BG, arrowcolor=COLOR_PRIMARY_D,
                    gripcount=0, relief="flat")
        s.map("Slim.Vertical.TScrollbar",
              background=[("active", COLOR_BORDER)])

    def _center_window(self, w, h):
        self.root.update_idletasks()
        sw = self.root.winfo_screenwidth()
        sh = self.root.winfo_screenheight()
        x = max(0, (sw - w) // 2)
        y = max(0, (sh - h) // 2 - 20)
        self.root.geometry(f"{w}x{h}+{x}+{y}")

    # ---------------- UI ----------------
    def _build_ui(self):
        # ---------- Header (fixed) ----------
        header = tk.Frame(self.root, bg=COLOR_PRIMARY, height=64)
        header.pack(fill=tk.X)
        header.pack_propagate(False)
        tk.Label(header,
                 text="Cummins / Crane BOM & Actual Issue Tool",
                 bg=COLOR_PRIMARY, fg="white",
                 font=("Segoe UI", 15, "bold")
                 ).pack(side=tk.LEFT, padx=20)
        tk.Label(header,
                 text="生产BOM / 生产缺料 / 实际发料回写",
                 bg=COLOR_PRIMARY, fg="#E3F2FD",
                 font=("Segoe UI", 10)
                 ).pack(side=tk.LEFT)

        # ---------- Scrollable body ----------
        body_wrap = ttk.Frame(self.root, style="TFrame")
        body_wrap.pack(fill=tk.BOTH, expand=True)

        canvas = tk.Canvas(body_wrap, bg=COLOR_BG, highlightthickness=0)
        vscroll = ttk.Scrollbar(body_wrap, orient="vertical",
                                command=canvas.yview,
                                style="Slim.Vertical.TScrollbar")
        canvas.configure(yscrollcommand=vscroll.set)

        vscroll.pack(side=tk.RIGHT, fill=tk.Y)
        canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        body = ttk.Frame(canvas, style="TFrame", padding=(16, 12, 16, 12))
        canvas_window = canvas.create_window((0, 0), window=body, anchor="nw")

        def _on_body_configure(event):
            canvas.configure(scrollregion=canvas.bbox("all"))

        def _on_canvas_configure(event):
            canvas.itemconfig(canvas_window, width=event.width)

        body.bind("<Configure>", _on_body_configure)
        canvas.bind("<Configure>", _on_canvas_configure)

        def _on_mousewheel(event):
            canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        def _on_mousewheel_linux(event):
            canvas.yview_scroll(-1 if event.num == 4 else 1, "units")

        # 只在鼠标位于 canvas 上时生效，避免干扰 Text 控件
        canvas.bind("<MouseWheel>", _on_mousewheel)
        canvas.bind("<Button-4>", _on_mousewheel_linux)
        canvas.bind("<Button-5>", _on_mousewheel_linux)

        # 递归绑定到 body 内所有子控件（Text 除外）
        def _bind_wheel_recursive(widget):
            if isinstance(widget, tk.Text):
                return
            widget.bind("<MouseWheel>", _on_mousewheel, add="+")
            widget.bind("<Button-4>", _on_mousewheel_linux, add="+")
            widget.bind("<Button-5>", _on_mousewheel_linux, add="+")
            for child in widget.winfo_children():
                _bind_wheel_recursive(child)

        self._bind_wheel_recursive = _bind_wheel_recursive
        self._body_canvas = canvas
        self._body_frame = body

        # ---------- Configuration ----------
        cfg = ttk.LabelFrame(body,
                             text=" Configuration  配置 ",
                             style="Card.TLabelframe", padding=(14, 10))
        cfg.pack(fill=tk.X, pady=(0, 10))

        # WiFi row
        ttk.Label(cfg, text="WiFi 当前网络:",
                  background=COLOR_BG_ALT).grid(
            row=0, column=0, sticky=tk.W, padx=(0, 6), pady=(0, 6))
        self.wifi_var = tk.StringVar(value="(checking...)")
        ttk.Label(cfg, textvariable=self.wifi_var,
                  background=COLOR_BG_ALT,
                  foreground=COLOR_PRIMARY_D, font=FONT_BOLD).grid(
            row=0, column=1, sticky=tk.W, padx=(0, 6), pady=(0, 6))

        ttk.Button(cfg, text="Refresh 刷新",
                   style="Secondary.TButton",
                   command=self._refresh_wifi_label).grid(
            row=0, column=2, padx=(8, 6), pady=(0, 6))
        ttk.Button(cfg, text="To AOI 切到AOI",
                   style="Secondary.TButton",
                   command=self._switch_to_aoi).grid(
            row=0, column=3, padx=(0, 6), pady=(0, 6))
        ttk.Button(cfg, text="Restore 切回原网",
                   style="Secondary.TButton",
                   command=self._switch_to_original).grid(
            row=0, column=4, padx=(0, 6), pady=(0, 6))

        # Warehouse ID row
        ttk.Label(cfg, text="Production Warehouse ID  生产在库仓库ID:",
                  background=COLOR_BG_ALT).grid(
            row=1, column=0, sticky=tk.W, padx=(0, 6), pady=(6, 0))

        self.prod_wh_var = tk.StringVar(value=str(self.production_warehouse_id))
        ttk.Entry(cfg, textvariable=self.prod_wh_var, width=10).grid(
            row=1, column=1, sticky=tk.W, pady=(6, 0))

        ttk.Label(cfg,
                  text="(默认 39，生产在库数据来源 / default 39, source of production stock)",
                  style="Muted.TLabel").grid(
            row=1, column=2, columnspan=2, sticky=tk.W, padx=(8, 0), pady=(6, 0))

        ttk.Button(cfg, text="Apply & Save  应用并保存",
                   style="Secondary.TButton",
                   command=self.on_apply_warehouse_id).grid(
            row=1, column=4, sticky=tk.W, padx=(8, 0), pady=(6, 0))

        # ---------- File drop ----------
        drop_frame = ttk.LabelFrame(body,
                                    text=" 1. BOM Excel File  选择BOM文件 ",
                                    style="Card.TLabelframe", padding=(14, 10))
        drop_frame.pack(fill=tk.X, pady=(0, 10))

        self.drop_label = tk.Label(
            drop_frame,
            text="Drag & drop the BOM .xlsx here\n"
                 "拖拽 BOM .xlsx 文件到这里\n\n"
                 "(or click Browse 或点击 Browse)",
            bg=COLOR_BG_ALT,
            fg=COLOR_MUTED,
            height=4,
            relief="solid",
            borderwidth=1,
            font=("Segoe UI", 11),
        )
        self.drop_label.pack(fill=tk.X, padx=8, pady=8)

        self.drop_label.drop_target_register(DND_FILES)
        self.drop_label.dnd_bind("<<Drop>>", self.on_drop)

        btn_frame = ttk.Frame(drop_frame, style="TFrame")
        btn_frame.pack(fill=tk.X, padx=8, pady=(0, 4))
        ttk.Button(btn_frame, text="Browse...  浏览",
                   style="Secondary.TButton",
                   command=self.browse_file).pack(side="left")
        ttk.Button(btn_frame, text="Clear  清除",
                   style="Secondary.TButton",
                   command=self.clear_file).pack(side="left", padx=6)

        # ---------- Actions ----------
        action_frame = ttk.LabelFrame(body,
                                      text=" 2. Actions  操作 ",
                                      style="Card.TLabelframe", padding=(14, 10))
        action_frame.pack(fill=tk.X, pady=(0, 10))

        ttk.Label(
            action_frame,
            text="⚠ 请选择一种 BOM 导出模式（二选一） / "
                 "Choose one BOM export mode (mutually exclusive):",
            style="Warn.TLabel"
        ).pack(anchor="w", pady=(0, 6))

        self.mode_var = tk.StringVar(value="issue")

        ttk.Radiobutton(
            action_frame,
            text="(1) Generate Production BOM  生成生产BOM（全部）",
            variable=self.mode_var, value="full"
        ).pack(anchor="w", padx=(8, 0))
        ttk.Label(
            action_frame,
            text="      → 输出全部物料，不做任何在库扣减 / "
                 "all items, no stock deduction",
            style="Muted.TLabel"
        ).pack(anchor="w", padx=(24, 0), pady=(0, 4))

        ttk.Radiobutton(
            action_frame,
            text="(2) Generate Actual Issue  生产缺料清单（扣生产在库）",
            variable=self.mode_var, value="issue"
        ).pack(anchor="w", padx=(8, 0))
        ttk.Label(
            action_frame,
            text="      → 需求 − 生产在库 = 需发料给生产的数量 / "
                 "demand − production stock",
            style="Muted.TLabel"
        ).pack(anchor="w", padx=(24, 0), pady=(0, 10))

        row_exec = ttk.Frame(action_frame, style="TFrame")
        row_exec.pack(fill=tk.X, pady=(0, 8))

        ttk.Button(row_exec, text="Generate  生成",
                   style="Primary.TButton",
                   command=self.on_generate).pack(side="left", padx=(0, 8))
        ttk.Button(row_exec,
                   text="(3) Write Back Actual Issued  写回实际发料",
                   style="Primary.TButton",
                   command=self.write_back_issued).pack(side="left")

        # ---------- Write-back rules (inline, bilingual) ----------
        rules_frame = ttk.LabelFrame(
            action_frame,
            text=" Write-back Rules  回写规则 ",
            style="Card.TLabelframe", padding=(10, 8))
        rules_frame.pack(fill=tk.X, pady=(4, 0))

        ttk.Label(
            rules_frame,
            text="E 列 / Column E  (实际发料 / Actual issued):",
            background=COLOR_BG_ALT, foreground=COLOR_PRIMARY_D,
            font=FONT_BOLD
        ).pack(anchor="w", pady=(0, 2))

        self._add_rule_line(rules_frame,
            "① 生产在库 + 实际发料 两者都有 → 相加，标浅蓝",
            "   Both production stock & actual issued > 0  → sum, LIGHT BLUE")

        self._add_rule_line(rules_frame,
            "② 只有实际发料 → 填实际发料，标绿",
            "   Only actual issued > 0  → actual issued, GREEN")

        self._add_rule_line(rules_frame,
            "③ 只有生产在库 → 填生产在库，不标色",
            "   Only production stock > 0  → production stock, no fill")

        self._add_rule_line(rules_frame,
            "④ 两者都没有 → 留空",
            "   Neither  → leave blank")

        ttk.Separator(rules_frame, orient="horizontal").pack(
            fill=tk.X, pady=(6, 4))

        ttk.Label(
            rules_frame,
            text="F 列 / Column F  (偏差 / Deviation) = E − D，负数标淡红",
            background=COLOR_BG_ALT, foreground=COLOR_PRIMARY_D,
            font=FONT_BOLD
        ).pack(anchor="w")
        ttk.Label(
            rules_frame,
            text="F = E − D ; if F < 0  →  F cell marked LIGHT RED",
            background=COLOR_BG_ALT, foreground=COLOR_MUTED,
            font=FONT_SUB
        ).pack(anchor="w")

        # ---------- Log ----------
        log_frame = ttk.LabelFrame(body,
                                   text=" 3. Log  日志 ",
                                   style="Card.TLabelframe", padding=(14, 10))
        log_frame.pack(fill=tk.X, pady=(0, 10))

        log_wrap = ttk.Frame(log_frame, style="TFrame")
        log_wrap.pack(fill=tk.BOTH, expand=True)

        self.log_text = tk.Text(
            log_wrap, height=10, wrap="word", font=FONT_MONO,
            bg=COLOR_BG, fg=COLOR_TEXT, relief="flat",
            highlightthickness=0, insertbackground=COLOR_PRIMARY
        )
        self.log_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        scrollbar = ttk.Scrollbar(log_wrap, orient="vertical",
                                  command=self.log_text.yview)
        scrollbar.pack(side=tk.RIGHT, fill="y")
        self.log_text.configure(yscrollcommand=scrollbar.set)
        self.log_text.tag_config("INFO", foreground=COLOR_TEXT)
        self.log_text.tag_config("WARN", foreground=COLOR_WARN)
        self.log_text.tag_config("ERROR", foreground=COLOR_ERROR)
        self.log_text.tag_config("OK", foreground=COLOR_SUCCESS)
        self.log_text.configure(state=tk.DISABLED)

        # ---------- Status bar (fixed at bottom) ----------
        status_wrap = tk.Frame(self.root, bg=COLOR_BG_ALT,
                               highlightbackground=COLOR_BORDER,
                               highlightthickness=1)
        status_wrap.pack(fill=tk.X, side=tk.BOTTOM)
        ttk.Label(status_wrap, textvariable=self.status_var,
                  style="Status.TLabel", padding=(10, 6)).pack(anchor=tk.W)

        # 递归绑定鼠标滚轮到所有子控件（在全部 UI 创建完之后）
        self.root.after(100, lambda: self._bind_wheel_recursive(body))

    def _add_rule_line(self, parent, zh, en):
        """添加中英双语的一行规则 / Add a bilingual rule line."""
        line = ttk.Frame(parent, style="TFrame")
        line.pack(fill=tk.X, pady=(1, 0))
        ttk.Label(line, text=zh,
                  background=COLOR_BG_ALT, foreground=COLOR_TEXT,
                  font=FONT_BASE).pack(anchor="w")
        ttk.Label(line, text=en,
                  background=COLOR_BG_ALT, foreground=COLOR_MUTED,
                  font=FONT_SUB).pack(anchor="w")

    # ---------------- Log ----------------
    def log(self, msg, level="INFO"):
        def _do():
            ts = datetime.now().strftime("%H:%M:%S")
            tag = {"INFO": "INFO", "WARN": "WARN",
                   "ERROR": "ERROR", "OK": "OK"}.get(level, "INFO")
            self.log_text.configure(state=tk.NORMAL)
            self.log_text.insert("end", f"[{ts}] {level}: {msg}\n", (tag,))
            self.log_text.see("end")
            self.log_text.configure(state=tk.DISABLED)
            self.root.update_idletasks()
        self.root.after(0, _do)

    # ---------------- Warehouse ID config ----------------
    def on_apply_warehouse_id(self):
        try:
            prod_id = parse_single_warehouse_id(self.prod_wh_var.get())
        except ValueError as e:
            messagebox.showerror("Invalid  无效", str(e))
            return

        changed = (prod_id != self.production_warehouse_id)
        self.production_warehouse_id = prod_id
        self.config_data["production_warehouse_id"] = prod_id

        if save_config(self.config_data):
            self.log(f"Config saved: production_warehouse_id={prod_id}", "OK")
        else:
            self.log("Config save failed / 配置保存失败", "ERROR")

        self.prod_wh_var.set(str(prod_id))

        if changed:
            self.status_var.set(
                f"Production warehouse ID saved 已保存: {prod_id}")
        else:
            self.status_var.set(
                f"Production warehouse ID unchanged 未改变: {prod_id}")

    # ---------------- WiFi ----------------
    def _refresh_wifi_label(self):
        ssid = WifiManager.get_current_ssid() or "(not connected 未连接)"
        self.wifi_var.set(ssid)

    def _switch_to_aoi(self):
        self.log(f"Switching WiFi to {WIFI_PROFILE_NAME}... 切换中...")
        if not WifiManager.ensure_profile():
            self.log(f"Profile {WIFI_PROFILE_NAME} missing 无法创建", "ERROR")
            messagebox.showerror("WiFi", f"Cannot ensure profile {WIFI_PROFILE_NAME}")
            return

        cur = WifiManager.get_current_ssid()
        if cur and cur != WIFI_PROFILE_NAME:
            self._original_wifi = cur

        def worker():
            ok = WifiManager.connect(WIFI_PROFILE_NAME)
            self.log("WiFi OK 已连接 AOI" if ok else "WiFi 连接失败",
                     "OK" if ok else "ERROR")
            self.root.after(0, self._refresh_wifi_label)
        threading.Thread(target=worker, daemon=True).start()

    def _switch_to_original(self):
        if not self._original_wifi:
            messagebox.showinfo("WiFi", "No original WiFi recorded 未记录原网络")
            return
        target = self._original_wifi
        self.log(f"Switching WiFi back to {target}... 切回原网...")

        def worker():
            ok = WifiManager.connect(target)
            self.log(f"WiFi switched back to {target}" if ok
                     else f"Failed to switch back to {target}",
                     "OK" if ok else "ERROR")
            self.root.after(0, self._refresh_wifi_label)
        threading.Thread(target=worker, daemon=True).start()

    # ---------------- File handling ----------------
    def on_drop(self, event):
        files = self.root.tk.splitlist(event.data)
        if not files:
            return
        path = files[0]
        if not path.lower().endswith((".xlsx", ".xlsm", ".xls")):
            messagebox.showwarning(
                "Invalid file  无效文件",
                "Please drop an Excel file (.xlsx).\n请拖拽 Excel 文件 (.xlsx)。")
            return
        self._load_file(path)

    def browse_file(self):
        path = filedialog.askopenfilename(
            title="Select BOM Excel  选择 BOM Excel",
            filetypes=[("Excel files", "*.xlsx *.xlsm *.xls"),
                       ("All files", "*.*")],
        )
        if path:
            self._load_file(path)

    def _load_file(self, path):
        self.file_path.set(path)
        self.base_name = os.path.splitext(os.path.basename(path))[0]
        self.bom_items = []
        self.bom_title = ""
        self.status_var.set(f"Loaded 已加载: {os.path.basename(path)}")
        self.log(f"File selected / 已选择文件: {path}")
        self.log(f"Base name / 基础文件名: {self.base_name}")

    def clear_file(self):
        self.file_path.set("")
        self.base_name = ""
        self.bom_items = []
        self.bom_title = ""
        self.production_stock = {}
        self.status_var.set(
            "Drag & drop an Excel file here, or click Browse.  "
            "拖拽 Excel 文件到此处，或点击 Browse。")
        self.log("File cleared / 文件已清除。")

    def _check_file(self):
        path = self.file_path.get().strip()
        if not path:
            messagebox.showwarning(
                "No file  无文件",
                "Please select or drop a BOM Excel file first.\n"
                "请先选择或拖入 BOM Excel 文件。")
            return None
        if not os.path.exists(path):
            messagebox.showerror(
                "File not found  文件不存在",
                f"File does not exist:\n{path}\n\n文件不存在。")
            return None
        return path

    # ---------------- Unified Generate ----------------
    def on_generate(self):
        if self.mode_var.get() == "full":
            self.generate_bom()
        else:
            self.generate_actual_issue()

    # ---------------- Mode 1 ----------------
    def generate_bom(self):
        path = self._check_file()
        if not path:
            return

        try:
            self.log("=" * 60)
            self.log("Mode (1): Generate Production BOM (ALL, no stock deduction)")
            self.log("模式 (1): 生成生产BOM（全部，不扣生产在库）")
            self.log("Reading BOM Excel ... / 正在读取 BOM ...")
            bom_title, items, _, _, _ = read_bom_excel(path)
            self.bom_title = bom_title
            self.bom_items = items

            self.log(f"BOM title: {bom_title}")
            self.log(f"Read {len(items)} BOM items. / 读取 {len(items)} 条物料。")

            rows = generate_bom_table(items)

            if not self.base_name:
                self.base_name = os.path.splitext(os.path.basename(path))[0]
            out_name = f"{self.base_name}_BOM.xlsx"
            out_dir = os.path.dirname(path)
            out_path = os.path.join(out_dir, out_name)

            export_bom_excel(rows, out_path)
            self.log(f"Production BOM exported: {out_path}", "OK")
            self.status_var.set(f"Production BOM exported 已导出: {out_name}")
            messagebox.showinfo(
                "Success  成功",
                f"生产BOM（全部需要发送的物料）已生成:\n"
                f"Production BOM (ALL items) generated:\n\n{out_path}\n\n"
                f"注意：这是所有需要发送的生产 BOM，未做任何库存扣减。\n"
                f"Note: full BOM, no stock deduction applied."
            )

        except Exception as e:
            self.log("ERROR: " + str(e), "ERROR")
            self.log(traceback.format_exc(), "ERROR")
            messagebox.showerror(
                "Error  错误",
                f"Failed to generate BOM:\n生成 BOM 失败:\n{e}")

    # ---------------- Mode 2 ----------------
    def generate_actual_issue(self):
        path = self._check_file()
        if not path:
            return

        if not self._sync_production_id_silent():
            return

        try:
            self.log("=" * 60)
            self.log("Mode (2): Generate Actual Issue (Production shortage only)")
            self.log("模式 (2): 生产缺料清单（扣生产在库）")
            self.log(f"Production warehouse ID = {self.production_warehouse_id}")

            self.log("Reading BOM Excel ... / 正在读取 BOM ...")
            bom_title, items, _, _, _ = read_bom_excel(path)
            self.bom_title = bom_title
            self.bom_items = items
            self.log(f"BOM title: {bom_title}")
            self.log(f"Read {len(items)} BOM items. / 读取 {len(items)} 条物料。")

            wh_id = self.production_warehouse_id
            self.log(f"Fetching production stock "
                     f"(warehouse_name_id={wh_id}) ...")
            self.log(f"正在查询生产在库 (warehouse_name_id={wh_id}) ...")
            warehouse_records = fetch_all_warehouse_by_id(wh_id, log_func=self.log)
            self.log(f"Fetched {len(warehouse_records)} production records.")

            production_stock = build_stock_map(warehouse_records)
            self.production_stock = production_stock
            self.log(f"Production stock part count: {len(production_stock)}")

            rows = generate_actual_issue_table(
                items, production_stock, log_func=self.log)
            self.log(f"Actual issue rows (only need issue): {len(rows)}")
            self.log(f"需要发料的行数: {len(rows)}")

            if not self.base_name:
                self.base_name = os.path.splitext(os.path.basename(path))[0]
            out_name = f"{self.base_name}_Actual_Issue.xlsx"
            out_dir = os.path.dirname(path)
            out_path = os.path.join(out_dir, out_name)

            export_actual_issue_excel(rows, out_path)
            self.log(f"Actual issue table exported: {out_path}", "OK")
            self.status_var.set(f"Actual issue exported 已导出: {out_name}")
            messagebox.showinfo(
                "Success  成功",
                f"生产缺料清单已生成:\n"
                f"Production actual issue generated:\n\n{out_path}\n\n"
                f"生产在库仓库ID / Production warehouse ID: {wh_id}\n"
                f"需要发料行数 / rows: {len(rows)}"
            )

        except Exception as e:
            self.log("ERROR: " + str(e), "ERROR")
            self.log(traceback.format_exc(), "ERROR")
            messagebox.showerror(
                "Error  错误",
                f"Failed to generate actual issue table:\n"
                f"生成缺料清单失败:\n{e}")

    def _sync_production_id_silent(self) -> bool:
        try:
            new_id = parse_single_warehouse_id(self.prod_wh_var.get())
        except ValueError:
            new_id = self.production_warehouse_id
        if new_id != self.production_warehouse_id:
            self.production_warehouse_id = new_id
            self.config_data["production_warehouse_id"] = new_id
            save_config(self.config_data)
            self.log(
                f"Production warehouse ID updated to {new_id} (auto-saved) / "
                f"生产在库仓库ID已更新为 {new_id}（已自动保存）", "OK")
        return True

    # ---------------- Mode 3 ----------------
    def write_back_issued(self):
        path = self._check_file()
        if not path:
            return

        # sync production warehouse id
        try:
            prod_id = parse_single_warehouse_id(self.prod_wh_var.get())
        except ValueError as e:
            messagebox.showerror("Invalid  无效", str(e))
            return
        if prod_id != self.production_warehouse_id:
            self.production_warehouse_id = prod_id
            self.config_data["production_warehouse_id"] = prod_id
            save_config(self.config_data)
            self.log(
                f"Production warehouse ID updated to {prod_id} (auto-saved) / "
                f"生产在库仓库ID已更新为 {prod_id}（已自动保存）", "OK")

        dlg = OrderIdDialog(self.root)
        order_id = dlg.result
        if not order_id:
            self.log("Write-back cancelled: no order ID. / 回写取消：未输入订单号。")
            return

        try:
            self.log("=" * 60)
            self.log("Mode (3): Write Back Actual Issued")
            self.log("模式 (3): 写回实际发料")

            # 1) 查生产在库
            self.log(f"Fetching production stock "
                     f"(warehouse_name_id={self.production_warehouse_id}) ...")
            self.log(f"正在查询生产在库 "
                     f"(warehouse_name_id={self.production_warehouse_id}) ...")
            warehouse_records = fetch_all_warehouse_by_id(
                self.production_warehouse_id, log_func=self.log)
            self.log(f"Fetched {len(warehouse_records)} warehouse records.")
            production_stock = build_stock_map(warehouse_records)
            self.production_stock = production_stock
            self.log(f"Production stock part count: {len(production_stock)}")

            # 2) 查订单
            self.log(f"Fetching order detail for: {order_id}")
            self.log(f"正在查询订单明细: {order_id}")
            items = fetch_order_detail(order_id)
            self.log(f"Order detail records: {len(items)}")

            issued_map = build_issued_map_from_order(items)
            self.log(f"Distinct part numbers with issued qty: {len(issued_map)}")

            # 3) 回写
            if not self.base_name:
                self.base_name = os.path.splitext(os.path.basename(path))[0]
            out_name = f"{self.base_name}_Actual_Issued_Written_Back.xlsx"
            out_dir = os.path.dirname(path)
            out_path = os.path.join(out_dir, out_name)

            updated = write_back_actual_issued(
                path, issued_map, production_stock, out_path, log_func=self.log
            )
            self.log(f"Write-back done. Filled E rows: {updated}", "OK")
            self.log(f"Output file: {out_path}", "OK")
            self.status_var.set(f"Write-back exported 已导出: {out_name}")
            messagebox.showinfo(
                "Success  成功",
                f"实际发料已回写:\n"
                f"Actual issued written back:\n\n{out_path}\n\n"
                f"E 列规则已在界面显示。\n"
                f"Write-back rules are shown on the main UI.\n"
                f"F 列 = E − D，负数标淡红。"
            )

        except Exception as e:
            self.log("ERROR: " + str(e), "ERROR")
            self.log(traceback.format_exc(), "ERROR")
            messagebox.showerror(
                "Error  错误",
                f"Failed to write back actual issued:\n"
                f"回写实际发料失败:\n{e}")


# =========================================================
# Entry
# =========================================================
def main():
    root = TkinterDnD.Tk()
    app = App(root)
    root.mainloop()


if __name__ == "__main__":
    main()