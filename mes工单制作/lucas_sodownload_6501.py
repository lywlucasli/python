import sys
import subprocess
import importlib.util
import os
from datetime import datetime

# ===================== 全局配置变量 =====================
GLOBAL_COMPANY_ID = "6501"
USER_NAME = "1441"
PASSWORD = "1442"
# -------------------------------------------------------

def auto_install(package_name: str):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package_name])

if not importlib.util.find_spec("requests"):
    print("[AUTO-INSTALL] requests not found, start installing...")
    auto_install("requests")
if not importlib.util.find_spec("xlwt"):
    print("[AUTO-INSTALL] xlwt not found, start installing...")
    auto_install("xlwt")

import hashlib
import requests
import xlwt
from typing import Optional, List, Dict

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("[AUTO-INSTALL] openpyxl not found, start installing...")
    auto_install("openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

try:
    import pyperclip
except ImportError:
    print("[AUTO-INSTALL] pyperclip not found, start installing...")
    auto_install("pyperclip")
    import pyperclip

# ===================== 辅助函数：获取西班牙语日期 =====================
def get_spanish_date_str(date_str: Optional[str]) -> str:
    if date_str and date_str.strip():
        return date_str.strip()
    now = datetime.now()
    months = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
              "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
    return now.strftime(f"%d-{months[now.month-1]}-%Y")

# ===================== 1. MES登录模块 =====================
class MES:
    def encrypt_password(self, password: str) -> str:
        md5 = hashlib.md5()
        md5.update(password.encode('utf-8'))
        return md5.hexdigest()

    def login(self, account: str, password: str) -> Optional[str]:
        print("[MES] Start login procedure...")
        encrypted_pwd = self.encrypt_password(password)
        url = "http://10.97.245.205:86/login"
        headers = {
            "Content-Type": "application/json;charset=UTF-8",
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
        }
        payload = {"account": account, "password": encrypted_pwd}
        try:
            resp = requests.post(url, headers=headers, json=payload, timeout=15)
            resp.raise_for_status()
            res_json = resp.json()
            if res_json.get("code") == 200 and "data" in res_json and "token" in res_json["data"]:
                print("[MES] Login success, token obtained")
                return res_json["data"]["token"]
            else:
                print(f"[MES] Login failed: {res_json.get('info', 'Unknown error')}")
                return None
        except Exception as e:
            print(f"[MES] Login exception: {str(e)}")
            return None

# ===================== 2. 剪贴板解析模块 =====================
def is_station_field(s: str) -> bool:
    return "_" in s and "/" not in s

def parse_clip_text(raw_text: str) -> List[Dict]:
    print("[PARSER] Start parsing clip text data (support no-date continuation lines)")
    result_list = []
    lines = raw_text.splitlines()
    last_station = None
    last_date = None

    for line_idx, raw_line in enumerate(lines, start=1):
        line = raw_line.strip()
        if not line:
            continue

        # 按制表符或空格分割
        if '\t' in line:
            parts = line.split('\t')
        else:
            parts = line.split()
        parts = [p for p in parts if p]

        if len(parts) < 4:   # 最少需要 WO, PN, 描述, Qty
            print(f"[PARSER] Skip line {line_idx}: too few fields ({len(parts)})")
            continue

        # 判断是否为格式A（包含客户名和日期）
        first = parts[0]
        second = parts[1] if len(parts) > 1 else ''
        is_format_a = False
        if '/' not in first and not first.isdigit() and ('/' in second or second.isdigit()):
            is_format_a = True
        elif len(parts) >= 6 and '/' not in first and not first.isdigit():
            is_format_a = True

        if is_format_a:
            # 格式A：客户名, 日期, 工单, 料号, 描述..., 数量
            # 统一 station 大小写（首字母大写）
            station_val = first.title()   # <--- 修改点
            date_val = second
            last_station = station_val
            last_date = date_val
            wo_val = parts[2]
            pn_val = parts[3]
            desc_and_qty = parts[4:]   # 描述和数量
        else:
            # 格式B：工单, 料号, 描述..., 数量 (无客户名和日期)
            if last_station is None:
                print(f"[PARSER] Skip line {line_idx}: no previous station and not format A")
                continue
            station_val = last_station
            date_val = last_date   # 沿用上一个日期
            wo_val = parts[0]
            pn_val = parts[1]
            desc_and_qty = parts[2:]   # 描述和数量

        # 提取数量：从后往前找数字
        if not desc_and_qty:
            print(f"[PARSER] Skip line {line_idx}: missing description and qty")
            continue
        qty_idx = -1
        qty_val = ""
        for i in range(len(desc_and_qty)-1, -1, -1):
            if desc_and_qty[i].replace('.', '').isdigit():
                qty_idx = i
                qty_val = desc_and_qty[i]
                break
        if qty_idx == -1:
            print(f"[PARSER] Skip line {line_idx}: cannot find qty in {desc_and_qty}")
            continue

        # 描述 = 数量之前的所有字段合并
        desc_val = " ".join(desc_and_qty[:qty_idx])
        # 数量之后的字段（如有）处理为 PMC 和备注
        after_qty = desc_and_qty[qty_idx+1:]
        if not after_qty:
            pmc_val = ""
            note_val = ""
        elif after_qty[0].lower() == "please":
            pmc_val = ""
            note_val = " ".join(after_qty)
        else:
            pmc_val = after_qty[0]
            note_val = " ".join(after_qty[1:])

        row_dict = {
            "station": station_val,
            "issue_date": date_val,
            "work_order": wo_val,
            "part_number": pn_val,
            "description": desc_val,
            "qty": qty_val,
            "pmc": pmc_val,
            "ready_note": note_val
        }
        result_list.append(row_dict)
        print(f"[PARSER] Line {line_idx} ok | station:{station_val} WO:{wo_val} PN:{pn_val} Qty:{qty_val} pmc:{repr(pmc_val)}")

    print(f"[PARSER] Parse complete, total records: {len(result_list)}")
    return result_list

# ===================== 3. 去重模块 =====================
def deduplicate_records(records: List[Dict]) -> tuple[List[Dict], List[Dict]]:
    print("[DEDUPLICATE] Start deduplication check")
    seen_keys = set()
    valid = []
    duplicate = []
    for rec in records:
        key = (rec["work_order"], str(rec["qty"]).strip())
        if key not in seen_keys:
            seen_keys.add(key)
            valid.append(rec)
        else:
            duplicate.append(rec)
    print(f"[DEDUPLICATE] Valid: {len(valid)} | Duplicate: {len(duplicate)}")
    return valid, duplicate

# ===================== 4. Excel美化工具（基础样式） =====================
HEADER_FILL = PatternFill("solid", fgColor="0070C0")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
ZEBRA_FILL = PatternFill("solid", fgColor="EBF1F8")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9")
)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT_ALIGN = Alignment(horizontal="right", vertical="center", wrap_text=True)
HIGHLIGHT_FILL = PatternFill("solid", fgColor="FFFF99")

def style_worksheet(ws, col_widths: List[int]):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 30
    for idx, w in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64+idx)].width = w
    for row_idx in range(2, ws.max_row + 1):
        fill = ZEBRA_FILL if (row_idx - 2) % 2 == 0 else None
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = THIN_BORDER
            if fill:
                cell.fill = fill
            if col_idx <= 2:
                cell.alignment = LEFT_ALIGN
            else:
                cell.alignment = RIGHT_ALIGN
    ws.freeze_panes = "A2"

def bin_to_multiline(bin_str: str) -> str:
    if not bin_str:
        return ""
    return bin_str.replace(",", "\n").replace(",", "\n")

# ===================== 5. 原始数据Excel写入 =====================
def save_two_sheet_excel(valid_rows: List[Dict], dup_rows: List[Dict], filepath: str):
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Valid_Records"
    headers = list(valid_rows[0].keys()) if valid_rows else []
    ws1.append(headers)
    for r in valid_rows:
        ws1.append([r.get(h,"") for h in headers])
    style_worksheet(ws1, [14,12,16,24,40,8,8,30])
    ws2 = wb.create_sheet("Duplicate_Records")
    dup_headers = list(dup_rows[0].keys()) if dup_rows else []
    ws2.append(dup_headers)
    for r in dup_rows:
        ws2.append([r.get(h,"") for h in dup_headers])
    style_worksheet(ws2, [14,12,16,24,40,8,8,30])
    wb.save(filepath)
    print(f"[EXCEL] Saved -> {filepath}")

def save_zmf60a(valid_rows: List[Dict], filepath: str):
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("ZMF60A")
    ws.write(0, 0, "plant")
    ws.write(0, 1, "SO")
    seen_wo = set()
    row = 1
    plant_num = int(GLOBAL_COMPANY_ID)
    for rec in valid_rows:
        wo = rec["work_order"]
        if wo not in seen_wo:
            seen_wo.add(wo)
            try:
                wo_num = int(wo)
            except ValueError:
                wo_num = wo
            ws.write(row, 0, plant_num)
            ws.write(row, 1, wo_num)
            row += 1
    wb.save(filepath)
    print(f"[ZMF60A] Saved -> {filepath}, unique WO: {len(seen_wo)}")

# ===================== 6. 按Station和料号前缀分组（已扩展55、57、59） =====================
def group_by_station_and_prefix(records: List[Dict]) -> Dict[str, Dict[str, List[Dict]]]:
    print("[GROUP] Start grouping")
    station_group: Dict[str, List[Dict]] = {}
    for rec in records:
        station = rec["station"]
        if station not in station_group:
            station_group[station] = []
        station_group[station].append(rec)

    final_group: Dict[str, Dict[str, List[Dict]]] = {}
    for station, station_records in station_group.items():
        g50, g52, g55_57_59, g8085 = [], [], [], []   # 合并 55/57/59
        for rec in station_records:
            pn = rec["part_number"]
            if pn.startswith("50"):
                g50.append(rec)
            elif pn.startswith("52"):
                g52.append(rec)
            elif pn.startswith(("55", "57", "59")):   # 三者合并
                g55_57_59.append(rec)
            elif pn.startswith("80") or pn.startswith("85"):
                g8085.append(rec)
        final_group[station] = {
            "prefix_50": g50,
            "prefix_52": g52,
            "prefix_55_57_59": g55_57_59,   # 合并后的组
            "prefix_80_85": g8085,
        }
        print(f"[GROUP] {station} | 50:{len(g50)} 52:{len(g52)} 55/57/59:{len(g55_57_59)} 80/85:{len(g8085)}")
    return final_group

# ===================== BOM接口（添加工单号） =====================
def fetch_work_order_bom(token: str, work_order: str) -> List[Dict]:
    url = (
        f"http://10.97.245.205:92/prod_baseapi/mes/wobomdetail/page"
        f"?rows=10000&page=1"
        f"&sidx=partno_id+asc,WORK_ORDER+DESC,rspos+asc"
        f"&sord="
        f"&company_id={GLOBAL_COMPANY_ID}"
        f"&work_order={work_order}"
        f"&lead_work_order="
        f"&partno_id="
    )
    headers = {
        "token": token,
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
    }
    out_rows = []
    try:
        resp = requests.get(url, headers=headers, timeout=20)
        resp.raise_for_status()
        js = resp.json()
        if js.get("code") != 200:
            print(f"[BOM-API] WO:{work_order} error: {js.get('info')}")
            return []
        rows = js.get("data", {}).get("rows", [])
        for row in rows:
            partno_id = row.get("partno_id", "").strip()
            req_qty = float(row.get("request_qty", 0.0))
            sobkz = row.get("sobkz", "")
            dumps = row.get("dumps", "")
            dbskz = row.get("dbskz", "")
            if req_qty <= 0:
                continue
            if sobkz == "Q" and dumps == "X":
                continue
            if dbskz == "E":
                continue
            uom = row.get("uom", "EA")
            out_rows.append({
                "partno_id": partno_id,
                "qty": req_qty,
                "work_order": work_order,
                "uom": uom
            })
        print(f"[BOM-API] WO:{work_order} valid rows: {len(out_rows)}")
        return out_rows
    except Exception as e:
        print(f"[BOM-API] WO:{work_order} exception: {str(e)}")
        return []

# ===================== BOM处理工具 =====================
def calc_agg_dict(bom_raw: List[Dict]) -> Dict[str, float]:
    agg_dict: Dict[str, float] = {}
    for item in bom_raw:
        pn = item["partno_id"]
        q = item["qty"]
        agg_dict[pn] = agg_dict.get(pn, 0.0) + q
    return agg_dict

def mark_last_row_total(bom_raw: List[Dict], agg_dict: Dict[str, float]) -> List[Dict]:
    last_idx_map: Dict[str, int] = {}
    for idx, it in enumerate(bom_raw):
        last_idx_map[it["partno_id"]] = idx
    output = []
    for idx, it in enumerate(bom_raw):
        pn = it["partno_id"]
        row = dict(it)
        if idx == last_idx_map[pn]:
            row["total_qty"] = agg_dict[pn]
        else:
            row["total_qty"] = None
        output.append(row)
    return output

# ===================== 库存接口 =====================
def fetch_all_inventory(token: str) -> List[Dict]:
    url = (
        f"http://10.97.245.205:92/prod_baseapi/wms/sendproductionbill/getSock/page"
        f"?rows=5000000&page=1"
        f"&sidx=++part_no+DESC,vtech_lot+asc+"
        f"&part_no="
        f"&bin="
        f"&company_id={GLOBAL_COMPANY_ID}"
        f"&statusList="
        f"&batch="
        f"&mgl=false"
        f"&uskey=false"
        f"&RMBPurchaseKey=false"
    )
    headers = {
        "token": token,
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
    }
    output = []
    try:
        print("[INVENTORY-API] Start downloading full inventory...")
        resp = requests.get(url, headers=headers, timeout=120)
        resp.raise_for_status()
        js = resp.json()
        if js.get("code") != 200:
            print(f"[INVENTORY-API] error: {js.get('info')}")
            return []
        rows = js.get("data", {}).get("rows", [])
        for row in rows:
            output.append({
                "bin_id": row.get("bin_id", "").strip(),
                "part_no": row.get("part_no", "").strip(),
                "store_qty": float(row.get("store_qty", 0.0)),
                "q_qty": float(row.get("q_qty", 0.0))
            })
        print(f"[INVENTORY-API] Download complete, rows: {len(output)}")
        return output
    except Exception as e:
        print(f"[INVENTORY-API] exception: {str(e)}")
        return []

def process_inventory(raw_inv: List[Dict]) -> Dict[str, Dict]:
    print("[INVENTORY-PROC] Start processing")
    inv_group: Dict[str, List[Dict]] = {}
    for rec in raw_inv:
        pn = rec["part_no"].strip()
        if pn not in inv_group:
            inv_group[pn] = []
        inv_group[pn].append(rec)

    inv_dict_out = {}
    for pn, group_rows in inv_group.items():
        total_stock = 0.0
        normal_bin_set = set()
        for row in group_rows:
            binid = row["bin_id"].strip()
            if not (binid.startswith("R") or binid.startswith("Q") or binid.startswith("T")):
                if binid:
                    normal_bin_set.add(binid)
            total_stock += (row["store_qty"] + row["q_qty"])
        bin_id_str = "\n".join(sorted(list(normal_bin_set)))
        inv_dict_out[pn] = {"bin_id": bin_id_str, "total_stock": total_stock}
    print(f"[INVENTORY-PROC] Complete, valid parts: {len(inv_dict_out)}")
    return inv_dict_out

# 文件名格式化
def format_wo_name(wo_list: List[str]) -> str:
    total_wo = len(wo_list)
    take = wo_list[:6]
    parts = []
    prefix = ""
    for wo in take:
        if len(wo)>=7:
            if not prefix:
                prefix = wo[:3]
            parts.append(wo[-4:])
    joined = "-".join([prefix] + parts)
    joined += f"_WO{total_wo}"
    if len(wo_list) > 6:
        joined += "_more"
    max_filename_len = 120
    if len(joined) > max_filename_len:
        joined = joined[:max_filename_len]
    return joined

# ===================== 核心生成函数（50/55/57/59/80/85通用） =====================
def generate_formatted_excel(records, bom_rows, agg_dict, inv_dict, station, prefix, filepath):
    customer = station.split('_')[0] if '_' in station else station
    wo_set = sorted(list({r["work_order"] for r in records}))
    summary_rows = []
    for idx, wo in enumerate(wo_set, 1):
        rec = next(r for r in records if r["work_order"] == wo)
        raw_note = rec.get("ready_note", "")
        if raw_note.startswith("Please ready materials on "):
            remark = raw_note.replace("Please ready materials on ", "")
        else:
            remark = raw_note
        summary_rows.append({
            "no": idx,
            "so": wo,
            "customer": station,
            "model": rec["part_number"],
            "qty": rec["qty"],
            "date": rec["issue_date"],
            "remark": remark
        })

    bom_sorted = sorted(bom_rows, key=lambda x: x["partno_id"])
    detail_rows = []
    for idx, item in enumerate(bom_sorted, 1):
        pn = item["partno_id"]
        inv_info = inv_dict.get(pn, {"bin_id": "", "total_stock": 0.0})
        detail_rows.append({
            "no": idx,
            "wo": item["work_order"],
            "pn": pn,
            "req_qty": item["qty"],
            "unit": item.get("uom", "EA"),
            "bin": inv_info["bin_id"],
            "stock": inv_info["total_stock"],
            "remark": "",
        })

    pn_groups = {}
    for i, row in enumerate(detail_rows):
        pn = row["pn"]
        if pn not in pn_groups:
            pn_groups[pn] = {"first_idx": i, "total_qty": agg_dict.get(pn, 0.0)}
        pn_groups[pn]["last_idx"] = i

    wb = Workbook()
    ws = wb.active
    ws.title = "Picking"
    col_widths = [27, 37, 52, 47, 51, 15, 35, 43, 36]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))

    cell_a1 = ws['A1']
    cell_a1.value = "Firma del obsequiador\nIssuer Signature\n执料人员签名"
    cell_a1.font = Font(name='Arial', bold=True, size=11)
    cell_a1.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 100

    ws.merge_cells('B1:H1')
    raw_date = records[0]["issue_date"] if records else None
    sample_date = get_spanish_date_str(raw_date)
    if prefix == "prefix_50":
        suffix = "50 Bill of Materials"
    elif prefix == "prefix_80_85":
        suffix = "80&85 Bill of Materials"
    elif prefix in ("prefix_55", "prefix_57", "prefix_59"):
        suffix = f"{prefix.replace('prefix_','')} Bill of Materials"
    elif prefix == "prefix_55_57_59":
        suffix = "55/57/59 Bill of Materials"
    else:
        suffix = "Bill of Materials"
        
    title = f"{customer} {sample_date} {suffix}"
    cell_title = ws['B1']
    cell_title.value = title
    cell_title.font = Font(name='Arial', bold=True, size=48)
    cell_title.alignment = Alignment(horizontal='center', vertical='center')

    cell_i1 = ws['I1']
    cell_i1.value = "Alimentación de materiales a tiempo\nMaterial feeding time\n送物料给SMT时间"
    cell_i1.font = Font(name='Arial', bold=True, size=11)
    cell_i1.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    summary_headers = ["NO.", "SO#", "Customer", "", "Model", "Qty", "", "Date", "Remark"]
    for col_idx, val in enumerate(summary_headers, 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.value = val
        font_color = "FF0000" if col_idx == 9 else "000000"
        cell.font = Font(name='Arial', bold=True, size=24, color=font_color)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws.row_dimensions[2].height = 50

    for r_idx, row_data in enumerate(summary_rows, start=3):
        vals = [row_data["no"], row_data["so"], row_data["customer"],
                "", row_data["model"], row_data["qty"], "",
                row_data["date"], row_data["remark"]]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=r_idx, column=col, value=val)
            cell.font = Font(name='Arial', bold=True, size=24)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border
        ws.row_dimensions[r_idx].height = 75

    detail_header_row = len(summary_rows) + 3
    detail_headers = [
        ("Número", "No.", "序号"),
        ("Orden de trabajo", "Work Order", "工单"),
        ("Número de pieza", "Part No.", "料号"),
        ("número de unidades", "Qty per WO", "需求数"),
        ("Cantidad demandada total", "Total Qty", "总需求数"),
        ("Unidad", "Unit", "单位"),
        ("Ubicación", "Location", "位置"),
        ("Cantidad en inventario", "Stock", "库存数"),
        ("Observación", "Remark", "备注"),
    ]
    for i, (es, en, zh) in enumerate(detail_headers):
        col = i + 1
        cell = ws.cell(row=detail_header_row, column=col)
        cell.value = f"{es}\n{en}\n{zh}"
        cell.font = Font(name='Arial', bold=True, size=18)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[detail_header_row].height = 88

    data_start_row = detail_header_row + 1
    current_pn = None
    for r_idx, row_data in enumerate(detail_rows, start=data_start_row):
        pn = row_data["pn"]
        is_first = (pn != current_pn)
        if is_first:
            current_pn = pn
            total_val = agg_dict.get(pn, 0.0)
        else:
            total_val = ""

        ws.cell(row=r_idx, column=1, value=row_data["no"])
        ws.cell(row=r_idx, column=2, value=row_data["wo"])
        ws.cell(row=r_idx, column=3, value=row_data["pn"])
        ws.cell(row=r_idx, column=4, value=row_data["req_qty"])
        ws.cell(row=r_idx, column=5, value=total_val)
        ws.cell(row=r_idx, column=6, value=row_data["unit"])
        cell_g = ws.cell(row=r_idx, column=7, value=row_data["bin"])
        cell_g.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.cell(row=r_idx, column=8, value=row_data["stock"])
        ws.cell(row=r_idx, column=9, value=row_data["remark"])

        ws.row_dimensions[r_idx].height = 70
        is_highlight = pn_groups.get(pn, {}).get("last_idx", -1) > pn_groups.get(pn, {}).get("first_idx", -1)
        for col in range(1, 10):
            cell = ws.cell(row=r_idx, column=col)
            cell.font = Font(name='Arial', bold=True, size=24)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border
            if is_highlight:
                cell.fill = HIGHLIGHT_FILL
        if '\n' in (row_data["bin"] or ""):
            ws.row_dimensions[r_idx].height = None

    max_row = data_start_row + len(detail_rows) - 1
    ws.print_area = f"A1:{get_column_letter(9)}{max_row}"
    ws.oddFooter.center.text = "Page &[Page] of &[Pages]"
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = False
    ws.page_setup.orientation = 'portrait'
    ws.page_margins.left = 0
    ws.page_margins.right = 0
    ws.page_margins.top = 0
    ws.page_margins.bottom = 0
    ws.page_margins.header = 0
    ws.page_margins.footer = 0

    wb.save(filepath)
    print(f"[FORMATTED EXCEL] Saved -> {filepath}")

def generate_formatted_excel_52(records, bom_rows, agg_dict, inv_dict, station, prefix, filepath):
    customer = station.split('_')[0] if '_' in station else station
    wo_set = sorted(list({r["work_order"] for r in records}))
    summary_rows = []
    for idx, wo in enumerate(wo_set, 1):
        rec = next(r for r in records if r["work_order"] == wo)
        raw_note = rec.get("ready_note", "")
        if raw_note.startswith("Please ready materials on "):
            remark = raw_note.replace("Please ready materials on ", "")
        else:
            remark = raw_note
        summary_rows.append({
            "no": idx,
            "so": wo,
            "customer": station,
            "model": rec["part_number"],
            "qty": rec["qty"],
            "date": rec["issue_date"],
            "remark": remark
        })

    detail_dict = {}
    for item in bom_rows:
        pn = item["partno_id"]
        if pn not in detail_dict:
            detail_dict[pn] = {"total_qty": 0.0}
        detail_dict[pn]["total_qty"] += item["qty"]

    detail_rows = []
    for idx, (pn, data) in enumerate(sorted(detail_dict.items()), 1):
        inv_info = inv_dict.get(pn, {"bin_id": "", "total_stock": 0.0})
        stock = inv_info["total_stock"]
        total_qty = data["total_qty"]
        detail_rows.append({
            "no": idx,
            "pn": pn,
            "bin": inv_info["bin_id"],
            "total_qty": total_qty,
            "stock": stock,
        })

    wb = Workbook()
    ws = wb.active
    ws.title = "Picking"
    col_widths = [20, 44, 38, 44, 50, 20, 20, 36]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))

    cell_a1 = ws['A1']
    cell_a1.value = "Firma del obsequiador\nIssuer Signature\n执料人员签名"
    cell_a1.font = Font(name='Arial', bold=True, size=11)
    cell_a1.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 100

    ws.merge_cells('B1:G1')
    raw_date = records[0]["issue_date"] if records else None
    sample_date = get_spanish_date_str(raw_date)
    title = f"{customer} {sample_date} 52 SMT Bill of Materials"
    cell_title = ws['B1']
    cell_title.value = title
    cell_title.font = Font(name='Arial', bold=True, size=48)
    cell_title.alignment = Alignment(horizontal='center', vertical='center')

    cell_h1 = ws['H1']
    cell_h1.value = "Alimentación de materiales a tiempo\nMaterial feeding time\n送物料给SMT时间"
    cell_h1.font = Font(name='Arial', bold=True, size=11)
    cell_h1.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    summary_headers = ["NO.", "SO#", "Customer", "Model", "B/T", "Qty", "Date", "Remark"]
    for col_idx, val in enumerate(summary_headers, 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.value = val
        font_color = "FF0000" if col_idx == 8 else "000000"
        cell.font = Font(name='Arial', bold=True, size=24, color=font_color)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws.row_dimensions[2].height = 57

    for r_idx, row_data in enumerate(summary_rows, start=3):
        vals = [row_data["no"], row_data["so"], row_data["customer"],
                row_data["model"], "", row_data["qty"],
                row_data["date"], row_data["remark"]]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=r_idx, column=col, value=val)
            cell.font = Font(name='Arial', bold=True, size=24)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border
        ws.row_dimensions[r_idx].height = 75

    detail_header_row = len(summary_rows) + 3
    detail_headers_52 = [
        ("Número", "No.", "序号"),
        ("Número de pieza", "Part No.", "料号"),
        ("Ubicación", "Location", "位置"),
        ("Cantidad demandada total", "Total Qty", "总需求数"),
        ("实际发料数量", "Actual issued", "实际发料"),
        ("偏差", "Deviation", "偏差"),
        ("Inventario", "Stock", "库存总数"),
        ("Observación", "Remark", "备注")
    ]
    for i, (es, en, zh) in enumerate(detail_headers_52):
        col = i + 1
        cell = ws.cell(row=detail_header_row, column=col)
        cell.value = f"{es}\n{en}\n{zh}"
        cell.font = Font(name='Arial', bold=True, size=18)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[detail_header_row].height = 88

    data_start_row = detail_header_row + 1
    for r_idx, row_data in enumerate(detail_rows, start=data_start_row):
        ws.cell(row=r_idx, column=1, value=row_data["no"])
        ws.cell(row=r_idx, column=2, value=row_data["pn"])
        cell_c = ws.cell(row=r_idx, column=3, value=row_data["bin"])
        cell_c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.cell(row=r_idx, column=4, value=row_data["total_qty"])
        ws.cell(row=r_idx, column=5, value="")
        ws.cell(row=r_idx, column=6, value="")
        ws.cell(row=r_idx, column=7, value=row_data["stock"])
        ws.cell(row=r_idx, column=8, value="")

        ws.row_dimensions[r_idx].height = 70
        for col in range(1, 9):
            cell = ws.cell(row=r_idx, column=col)
            cell.font = Font(name='Arial', bold=True, size=24)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border
        if '\n' in (row_data["bin"] or ""):
            ws.row_dimensions[r_idx].height = None

    max_row = data_start_row + len(detail_rows) - 1
    ws.print_area = f"A1:{get_column_letter(8)}{max_row}"
    ws.oddFooter.center.text = "Page &[Page] of &[Pages]"
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = False
    ws.page_setup.orientation = 'portrait'
    ws.page_margins.left = 0
    ws.page_margins.right = 0
    ws.page_margins.top = 0
    ws.page_margins.bottom = 0
    ws.page_margins.header = 0
    ws.page_margins.footer = 0

    wb.save(filepath)
    print(f"[FORMATTED EXCEL 52] Saved -> {filepath}")

# ===================== 原简单导出函数（备用） =====================
def export_final_picking_sheet(rows: List[Dict], filepath: str, sheet_name: str = "Picking"):
    rows_sorted = sorted(rows, key=lambda x: x["pn"])
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(["PN", "Bin_Location", "Req_Qty", "Total_Qty", "Stock_Qty"])
    for r in rows_sorted:
        ws.append([
            r["pn"],
            bin_to_multiline(r.get("bin_id", "")),
            r.get("req_qty", ""),
            r.get("total_qty", ""),
            r.get("stock_qty", "")
        ])
    style_worksheet(ws, [24, 30, 14, 14, 14])
    wb.save(filepath)
    print(f"[FINAL-EXCEL] Saved -> {filepath} | rows:{len(rows_sorted)}")

# ===================== 主入口 =====================
if __name__ == "__main__":
    print("===== Material Processing Tool Start =====")
    print("[MAIN] Reading clipboard...")
    raw_clip = pyperclip.paste()
    if not raw_clip.strip():
        print("[MAIN ERROR] Clipboard empty!")
        sys.exit(1)

    parsed_all = parse_clip_text(raw_clip)
    if not parsed_all:
        print("[MAIN ERROR] No valid records")
        sys.exit(1)

    valid_records, dup_records = deduplicate_records(parsed_all)
    save_two_sheet_excel(valid_records, dup_records, "material_data.xlsx")
    save_zmf60a(valid_records, "ZMF60A.xls")

    zmf_path = os.path.abspath("ZMF60A.xls")
    pyperclip.copy(zmf_path)
    print(f"[CLIPBOARD] Copied ZMF60A.xls path: {zmf_path}")

    mes = MES()
    token = mes.login(USER_NAME, PASSWORD)
    if not token:
        print("[MAIN] MES login failed")
        sys.exit(1)

    station_group = group_by_station_and_prefix(valid_records)
    inv_dict = process_inventory(fetch_all_inventory(token))

    all_final_rows = []

    # ========== 修改后的主循环：遍历所有分组 ==========
    for station, group_data in station_group.items():
        print(f"\n[MAIN] ====== Station: {station} ======")
        for key, group in group_data.items():
            if not group:
                continue
            print(f"[MAIN] Processing {key} with {len(group)} records")
            # 跳过含 -P 的料号（与原有逻辑一致）
            filter_group = [rec for rec in group if "-P" not in rec["part_number"]]
            skip_count = len(group) - len(filter_group)
            if skip_count:
                print(f"[MAIN] Skip {skip_count} records (part contains -P, skip BOM download)")

            if not filter_group:
                continue

            wo_set = list({r["work_order"] for r in filter_group})
            bom_rows = []
            for wo in wo_set:
                bom_rows.extend(fetch_work_order_bom(token, wo))
            if not bom_rows:
                print(f"[MAIN WARNING] station={station} key={key} wo_set={wo_set} got zero bom rows!")
                continue

            agg_dict = calc_agg_dict(bom_rows)
            wo_fmt_str = format_wo_name(wo_set)
            out_file = f"{station}_{key}_{wo_fmt_str}.xlsx"

            if key == "prefix_52":
                generate_formatted_excel_52(filter_group, bom_rows, agg_dict, inv_dict, station, key, out_file)
            else:
                # 其他前缀（50、55、57、59、80_85）都走通用格式
                generate_formatted_excel(filter_group, bom_rows, agg_dict, inv_dict, station, key, out_file)

            # 同时收集到全局汇总（旧格式）
            marked_list = mark_last_row_total(bom_rows, agg_dict)
            for row in marked_list:
                pn = row["partno_id"]
                inv_info = inv_dict.get(pn, {"bin_id":"","total_stock":0.0})
                all_final_rows.append({
                    "pn": pn,
                    "bin_id": inv_info["bin_id"],
                    "req_qty": row["qty"],
                    "total_qty": row.get("total_qty"),
                    "stock_qty": inv_info["total_stock"]
                })

    if all_final_rows:
        export_final_picking_sheet(all_final_rows, "bom_stock_result.xlsx", sheet_name="All_Picking")

    print("\n" + "="*80)
    print("===== FINAL EXPORT PREVIEW (FIRST 15, SORTED BY PN) =====")
    print("="*80)
    sorted_preview = sorted(all_final_rows, key=lambda x: x["pn"])
    for idx, item in enumerate(sorted_preview[:15], 1):
        print(f"  {idx}. PN:{item['pn']} | Bin:{item['bin_id']} | Req:{item['req_qty']} | Total:{item['total_qty']} | Stock:{item['stock_qty']}")

    print("\n===== COMPLETED =====")
    print("Generated files:")
    print("  1. material_data.xlsx      - raw parse data")
    print("  2. ZMF60A.xls              - WO list (xls format, numeric cell)")
    print("  3. {station}_{prefix}_{wo}.xlsx - group picking sheet (all layers)")
    print("  4. bom_stock_result.xlsx   - global picking sheet (old format)")
    print(f"\nZMF60A.xls The path has been copied to the clipboard: {zmf_path}")
    input("\nPress Enter to exit...")