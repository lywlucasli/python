import sys
import subprocess
import importlib.util
import os
import threading
import tkinter as tk
import re
from tkinter import messagebox, ttk
from datetime import datetime

# ===================== 全局配置变量 =====================
GLOBAL_COMPANY_ID = "6501"
USER_NAME = "NAF684"
PASSWORD = "666888"
# -------------------------------------------------------

def auto_install(package_name: str):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package_name])

if not importlib.util.find_spec("requests"):
    print("[AUTO-INSTALL] requests not found, start installing...")
    auto_install("requests")
if not importlib.util.find_spec("xlwt"):
    print("[AUTO-INSTALL] xlwt not found, start installing...")
    auto_install("xlwt")

import hashlib
import requests
import xlwt
from typing import Optional, List, Dict

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("[AUTO-INSTALL] openpyxl not found, start installing...")
    auto_install("openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

try:
    import pyperclip
except ImportError:
    print("[AUTO-INSTALL] pyperclip not found, start installing...")
    auto_install("pyperclip")
    import pyperclip

try:
    from PIL import Image, ImageTk
except ImportError:
    print("[AUTO-INSTALL] Pillow not found, start installing...")
    auto_install("Pillow")
    from PIL import Image, ImageTk

# ===================== 辅助函数：获取西班牙语日期 =====================
def get_spanish_date_str(date_str: Optional[str]) -> str:
    if date_str and date_str.strip():
        return date_str.strip()
    now = datetime.now()
    months = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
              "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
    return now.strftime(f"%d-{months[now.month-1]}-%Y")

# ===================== 1. MES登录模块 =====================
class MES:
    def encrypt_password(self, password: str) -> str:
        md5 = hashlib.md5()
        md5.update(password.encode('utf-8'))
        return md5.hexdigest()

    def login(self, account: str, password: str) -> Optional[str]:
        print("[MES] Start login procedure...")
        encrypted_pwd = self.encrypt_password(password)
        url = "http://10.97.245.205:86/login"
        headers = {
            "Content-Type": "application/json;charset=UTF-8",
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
        }
        payload = {"account": account, "password": encrypted_pwd}
        try:
            resp = requests.post(url, headers=headers, json=payload, timeout=15)
            resp.raise_for_status()
            res_json = resp.json()
            if res_json.get("code") == 200 and "data" in res_json and "token" in res_json["data"]:
                print("[MES] Login success, token obtained")
                return res_json["data"]["token"]
            else:
                print(f"[MES] Login failed: {res_json.get('info', 'Unknown error')}")
                return None
        except Exception as e:
            print(f"[MES] Login exception: {str(e)}")
            return None

# ===================== 2. 剪贴板解析模块 =====================
def is_station_field(s: str) -> bool:
    return "_" in s and "/" not in s


def normalize_station(value: str) -> str:
    return " ".join(value.strip().split()).title()


def is_work_order(value: str) -> bool:
    value = value.strip()
    return value.isdigit() and len(value) in (8, 12)


def is_part_number(value: str) -> bool:
    value = value.strip()
    return len(value) == 18 and bool(re.fullmatch(r"[A-Za-z0-9-]+", value))


def is_quantity(value: str) -> bool:
    value = value.strip()
    try:
        return bool(value) and float(value) > 0
    except ValueError:
        return False


def is_issue_date(value: str) -> bool:
    value = value.strip()
    return bool(re.search(r"\d{1,4}[-/]\d{1,2}(?:[-/]\d{1,4})?", value))


def today_issue_date() -> str:
    return datetime.now().strftime("%m/%d/%Y")


def work_order_sort_key(work_order: str):
    value = str(work_order).strip()
    return (0, int(value)) if value.isdigit() else (1, value)

def parse_clip_text(raw_text: str) -> List[Dict]:
    print("[PARSER] Start parsing clip text data (support no-date continuation lines)")
    result_list = []
    lines = raw_text.splitlines()
    last_station = None
    last_date = None

    for line_idx, raw_line in enumerate(lines, start=1):
        line = raw_line.rstrip()
        if not line.strip():
            continue

        # 按制表符或空格分割
        has_tabs = '\t' in line
        if has_tabs:
            parts = [p.strip() for p in line.split('\t')]
        else:
            parts = line.strip().split()
        if not has_tabs:
            parts = [p for p in parts if p]

        # Excel复制内容通常包含表头，表头不是一条工单记录。
        header_text = " ".join(p for p in parts if p).lower()
        if ("customer" in header_text and ("so#" in header_text or "model" in header_text)):
            print(f"[PARSER] Skip header line {line_idx}")
            continue

        if len(parts) < 3:
            print(f"[PARSER] Skip line {line_idx}: too few fields ({len(parts)})")
            continue

        # 先找 18 位料号，再找它之前最近的 8/12 位工单号，避免列空值错位。
        pn_indexes = [idx for idx, value in enumerate(parts) if is_part_number(value)]
        if len(pn_indexes) != 1:
            print(f"[PARSER ERROR] Line {line_idx}: expected one 18-character part number, found {len(pn_indexes)}")
            continue
        pn_idx = pn_indexes[0]
        wo_indexes = [idx for idx, value in enumerate(parts[:pn_idx]) if is_work_order(value)]
        if not wo_indexes:
            print(f"[PARSER ERROR] Line {line_idx}: 8/12-digit work order not found before part number")
            continue
        wo_idx = wo_indexes[-1]
        if pn_idx != wo_idx + 1 and has_tabs:
            print(f"[PARSER ERROR] Line {line_idx}: work order and part number are not adjacent")
            continue

        # 根据工单号和料号定位客户、日期，空值沿用上一条或使用今天日期。
        station_parts = [value for value in parts[:wo_idx] if value and not is_issue_date(value)]
        station_val = normalize_station(" ".join(station_parts)) if station_parts else last_station
        date_candidates = [value for value in parts[:wo_idx] if is_issue_date(value)]
        date_val = date_candidates[0] if date_candidates else (last_date or today_issue_date())
        if station_val:
            last_station = station_val
        last_date = date_val
        wo_val = parts[wo_idx]
        pn_val = parts[pn_idx]
        after_part = parts[pn_idx + 1:]

        # 数量应在料号之后；优先取第一个正数，防止备注中的日期/数字误判。
        qty_idx = next((idx for idx, value in enumerate(after_part) if is_quantity(value)), -1)
        if qty_idx < 0:
            print(f"[PARSER ERROR] Line {line_idx}: quantity not found after part number")
            continue
        qty_val = after_part[qty_idx]

        # 数量之前的字段为描述；数量后的两列分别是 PMC 和备注，可为空。
        desc_val = " ".join(value for value in after_part[:qty_idx] if value)
        # 数量之后的字段（如有）处理为 PMC 和备注
        after_qty = after_part[qty_idx + 1:]
        if has_tabs and len(after_qty) > 2:
            print(f"[PARSER ERROR] Line {line_idx}: too many columns after quantity")
            continue
        pmc_val = after_qty[0] if after_qty else ""
        note_val = " ".join(after_qty[1:] if after_qty else [])
        if pmc_val.lower() == "please":
            note_val = " ".join(after_qty)
            pmc_val = ""

        row_dict = {
            "station": station_val,
            "issue_date": date_val,
            "work_order": wo_val,
            "part_number": pn_val,
            "description": desc_val,
            "qty": qty_val,
            "pmc": pmc_val,
            "ready_note": note_val
        }
        result_list.append(row_dict)
        print(f"[PARSER] Line {line_idx} ok | station:{station_val} WO:{wo_val} PN:{pn_val} Qty:{qty_val} pmc:{repr(pmc_val)}")

    print(f"[PARSER] Parse complete, total records: {len(result_list)}")
    return result_list

# ===================== 3. 去重模块 =====================
def deduplicate_records(records: List[Dict]) -> tuple[List[Dict], List[Dict]]:
    print("[DEDUPLICATE] Start deduplication check")
    seen_keys = set()
    valid = []
    duplicate = []
    for rec in records:
        key = (rec["work_order"], str(rec["qty"]).strip())
        if key not in seen_keys:
            seen_keys.add(key)
            valid.append(rec)
        else:
            duplicate.append(rec)
    print(f"[DEDUPLICATE] Valid: {len(valid)} | Duplicate: {len(duplicate)}")
    return valid, duplicate

# ===================== 4. Excel美化工具（基础样式） =====================
HEADER_FILL = PatternFill("solid", fgColor="0070C0")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
ZEBRA_FILL = PatternFill("solid", fgColor="EBF1F8")
THIN_BORDER = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000")
)
CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT_ALIGN = Alignment(horizontal="right", vertical="center", wrap_text=True)
HIGHLIGHT_FILL = PatternFill("solid", fgColor="FFFF99")

def style_worksheet(ws, col_widths: List[int]):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 30
    for idx, w in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64+idx)].width = w
    for row_idx in range(2, ws.max_row + 1):
        fill = ZEBRA_FILL if (row_idx - 2) % 2 == 0 else None
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = THIN_BORDER
            if fill:
                cell.fill = fill
            if col_idx == 3:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif col_idx <= 2:
                cell.alignment = LEFT_ALIGN
            else:
                cell.alignment = RIGHT_ALIGN
    ws.freeze_panes = "A2"

def bin_to_multiline(bin_str: str) -> str:
    if not bin_str:
        return ""
    return bin_str.replace(",", "\n").replace(",", "\n")

# ===================== 5. 原始数据Excel写入 =====================
def save_two_sheet_excel(valid_rows: List[Dict], dup_rows: List[Dict], filepath: str):
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Valid_Records"
    headers = list(valid_rows[0].keys()) if valid_rows else []
    ws1.append(headers)
    for r in valid_rows:
        ws1.append([r.get(h,"") for h in headers])
    style_worksheet(ws1, [14,12,16,24,40,8,8,30])
    ws2 = wb.create_sheet("Duplicate_Records")
    dup_headers = list(dup_rows[0].keys()) if dup_rows else []
    ws2.append(dup_headers)
    for r in dup_rows:
        ws2.append([r.get(h,"") for h in dup_headers])
    style_worksheet(ws2, [14,12,16,24,40,8,8,30])
    wb.save(filepath)
    print(f"[EXCEL] Saved -> {filepath}")

def save_zmf60a(valid_rows: List[Dict], filepath: str):
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("ZMF60A")
    ws.write(0, 0, "plant")
    ws.write(0, 1, "SO")
    seen_wo = set()
    row = 1
    plant_num = int(GLOBAL_COMPANY_ID)
    for rec in valid_rows:
        wo = rec["work_order"]
        if wo not in seen_wo:
            seen_wo.add(wo)
            try:
                wo_num = int(wo)
            except ValueError:
                wo_num = wo
            ws.write(row, 0, plant_num)
            ws.write(row, 1, wo_num)
            row += 1
    wb.save(filepath)
    print(f"[ZMF60A] Saved -> {filepath}, unique WO: {len(seen_wo)}")

# ===================== 6. 按Station和料号前缀分组（已扩展55、57、59） =====================
def group_by_station_and_prefix(records: List[Dict]) -> Dict[str, Dict[str, List[Dict]]]:
    print("[GROUP] Start grouping")
    station_group: Dict[str, List[Dict]] = {}
    for rec in records:
        station = rec["station"]
        if station not in station_group:
            station_group[station] = []
        station_group[station].append(rec)

    final_group: Dict[str, Dict[str, List[Dict]]] = {}
    for station, station_records in station_group.items():
        g50, g52, g55_57_59, g8085 = [], [], [], []
        for rec in station_records:
            pn = rec["part_number"]
            if pn.startswith("50"):
                g50.append(rec)
            elif pn.startswith("52"):
                g52.append(rec)
            elif pn.startswith(("55", "57", "59")):
                g55_57_59.append(rec)
            elif pn.startswith("80") or pn.startswith("85"):
                g8085.append(rec)
        final_group[station] = {
            "prefix_50": g50,
            "prefix_52": g52,
            "prefix_55_57_59": g55_57_59,
            "prefix_80_85": g8085,
        }
        print(f"[GROUP] {station} | 50:{len(g50)} 52:{len(g52)} 55/57/59:{len(g55_57_59)} 80/85:{len(g8085)}")
    return final_group

# ===================== BOM接口（添加工单号） =====================
def fetch_work_order_bom(token: str, work_order: str) -> List[Dict]:
    url = (
        f"http://10.97.245.205:92/prod_baseapi/mes/wobomdetail/page"
        f"?rows=10000&page=1"
        f"&sidx=partno_id+asc,WORK_ORDER+DESC,rspos+asc"
        f"&sord="
        f"&company_id={GLOBAL_COMPANY_ID}"
        f"&work_order={work_order}"
        f"&lead_work_order="
        f"&partno_id="
    )
    headers = {
        "token": token,
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
    }
    out_rows = []
    try:
        resp = requests.get(url, headers=headers, timeout=20)
        resp.raise_for_status()
        js = resp.json()
        if js.get("code") != 200:
            print(f"[BOM-API] WO:{work_order} error: {js.get('info')}")
            return []
        rows = js.get("data", {}).get("rows", [])
        for row in rows:
            partno_id = row.get("partno_id", "").strip()
            req_qty = float(row.get("request_qty", 0.0))
            sobkz = row.get("sobkz", "")
            dumps = row.get("dumps", "")
            dbskz = row.get("dbskz", "")
            if req_qty <= 0:
                continue
            if sobkz == "Q" and dumps == "X":
                continue
            if dbskz == "E":
                continue
            uom = row.get("uom", "EA")
            out_rows.append({
                "partno_id": partno_id,
                "qty": req_qty,
                "work_order": work_order,
                "uom": uom
            })
        print(f"[BOM-API] WO:{work_order} valid rows: {len(out_rows)}")
        return out_rows
    except Exception as e:
        print(f"[BOM-API] WO:{work_order} exception: {str(e)}")
        return []

# ===================== BOM处理工具 =====================
def calc_agg_dict(bom_raw: List[Dict]) -> Dict[str, float]:
    agg_dict: Dict[str, float] = {}
    for item in bom_raw:
        pn = item["partno_id"]
        q = item["qty"]
        agg_dict[pn] = agg_dict.get(pn, 0.0) + q
    return agg_dict

def mark_last_row_total(bom_raw: List[Dict], agg_dict: Dict[str, float]) -> List[Dict]:
    last_idx_map: Dict[str, int] = {}
    for idx, it in enumerate(bom_raw):
        last_idx_map[it["partno_id"]] = idx
    output = []
    for idx, it in enumerate(bom_raw):
        pn = it["partno_id"]
        row = dict(it)
        if idx == last_idx_map[pn]:
            row["total_qty"] = agg_dict[pn]
        else:
            row["total_qty"] = None
        output.append(row)
    return output

# ===================== 库存接口 =====================
def fetch_all_inventory(token: str) -> List[Dict]:
    url = (
        f"http://10.97.245.205:92/prod_baseapi/wms/sendproductionbill/getSock/page"
        f"?rows=5000000&page=1"
        f"&sidx=++part_no+DESC,vtech_lot+asc+"
        f"&part_no="
        f"&bin="
        f"&company_id={GLOBAL_COMPANY_ID}"
        f"&statusList="
        f"&batch="
        f"&mgl=false"
        f"&uskey=false"
        f"&RMBPurchaseKey=false"
    )
    headers = {
        "token": token,
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
    }
    output = []
    try:
        print("[INVENTORY-API] Start downloading full inventory...")
        resp = requests.get(url, headers=headers, timeout=120)
        resp.raise_for_status()
        js = resp.json()
        if js.get("code") != 200:
            print(f"[INVENTORY-API] error: {js.get('info')}")
            return []
        rows = js.get("data", {}).get("rows", [])
        for row in rows:
            output.append({
                "bin_id": row.get("bin_id", "").strip(),
                "part_no": row.get("part_no", "").strip(),
                "store_qty": float(row.get("store_qty", 0.0)),
                "q_qty": float(row.get("q_qty", 0.0))
            })
        print(f"[INVENTORY-API] Download complete, rows: {len(output)}")
        return output
    except Exception as e:
        print(f"[INVENTORY-API] exception: {str(e)}")
        return []

def process_inventory(raw_inv: List[Dict]) -> Dict[str, Dict]:
    print("[INVENTORY-PROC] Start processing")
    inv_group: Dict[str, List[Dict]] = {}
    for rec in raw_inv:
        pn = rec["part_no"].strip()
        if pn not in inv_group:
            inv_group[pn] = []
        inv_group[pn].append(rec)

    inv_dict_out = {}
    for pn, group_rows in inv_group.items():
        total_stock = 0.0
        normal_bin_set = set()
        for row in group_rows:
            binid = row["bin_id"].strip()
            if not (binid.startswith("R") or binid.startswith("Q") or binid.startswith("T")):
                if binid:
                    normal_bin_set.add(binid)
            total_stock += (row["store_qty"] + row["q_qty"])
        bin_id_str = "\n".join(sorted(list(normal_bin_set)))
        inv_dict_out[pn] = {"bin_id": bin_id_str, "total_stock": total_stock}
    print(f"[INVENTORY-PROC] Complete, valid parts: {len(inv_dict_out)}")
    return inv_dict_out

# 文件名格式化
def format_wo_name(wo_list: List[str]) -> str:
    total_wo = len(wo_list)
    take = wo_list[:6]
    parts = []
    prefix = ""
    for wo in take:
        if len(wo)>=7:
            if not prefix:
                prefix = wo[:3]
            parts.append(wo[-4:])
    joined = "-".join([prefix] + parts)
    joined += f"_WO{total_wo}"
    if len(wo_list) > 6:
        joined += "_more"
    max_filename_len = 120
    if len(joined) > max_filename_len:
        joined = joined[:max_filename_len]
    return joined

# ===================== 核心生成函数（50/55/57/59/80/85通用） =====================
def generate_formatted_excel(records, bom_rows, agg_dict, inv_dict, station, prefix, filepath):
    customer = station.split('_')[0] if '_' in station else station
    wo_set = sorted({r["work_order"] for r in records}, key=work_order_sort_key)
    summary_rows = []
    for idx, wo in enumerate(wo_set, 1):
        rec = next(r for r in records if r["work_order"] == wo)
        raw_note = rec.get("ready_note", "")
        if raw_note.startswith("Please ready materials on "):
            remark = raw_note.replace("Please ready materials on ", "")
        else:
            remark = raw_note
        summary_rows.append({
            "no": idx,
            "so": wo,
            "customer": station,
            "model": rec["part_number"],
            "qty": rec["qty"],
            "date": rec["issue_date"],
            "remark": remark
        })

    bom_sorted = sorted(bom_rows, key=lambda x: (x["partno_id"], work_order_sort_key(x["work_order"])))
    detail_rows = []
    for idx, item in enumerate(bom_sorted, 1):
        pn = item["partno_id"]
        inv_info = inv_dict.get(pn, {"bin_id": "", "total_stock": 0.0})
        detail_rows.append({
            "no": idx,
            "wo": item["work_order"],
            "pn": pn,
            "req_qty": item["qty"],
            "unit": item.get("uom", "EA"),
            "bin": inv_info["bin_id"],
            "stock": inv_info["total_stock"],
            "remark": "",
        })

    pn_groups = {}
    for i, row in enumerate(detail_rows):
        pn = row["pn"]
        if pn not in pn_groups:
            pn_groups[pn] = {"first_idx": i, "total_qty": agg_dict.get(pn, 0.0)}
        pn_groups[pn]["last_idx"] = i

    wb = Workbook()
    ws = wb.active
    ws.title = "Picking"
    col_widths = [27, 37, 52, 47, 51, 15, 35, 43, 36]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    thin_border = Border(left=Side(style='thin', color='000000'), right=Side(style='thin', color='000000'),
                         top=Side(style='thin', color='000000'), bottom=Side(style='thin', color='000000'))
    for row in range(1, 3):
        for col in range(1, 10):
            ws.cell(row=row, column=col).border = thin_border

    cell_a1 = ws['A1']
    cell_a1.value = "Firma del obsequiador\nIssuer Signature\n执料人员签名"
    cell_a1.font = Font(name='Arial', bold=True, size=11)
    cell_a1.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 100

    ws.merge_cells('B1:H2')
    ws['A2'] = None
    for row in range(1, 3):
        for col in range(2, 9):
            ws.cell(row=row, column=col).border = thin_border
    raw_date = records[0]["issue_date"] if records else None
    sample_date = get_spanish_date_str(raw_date)
    if prefix == "prefix_50":
        suffix = "50 Bill of Materials"
    elif prefix == "prefix_80_85":
        suffix = "80&85 Bill of Materials"
    elif prefix in ("prefix_55", "prefix_57", "prefix_59"):
        suffix = f"{prefix.replace('prefix_','')} Bill of Materials"
    elif prefix == "prefix_55_57_59":
        suffix = "55/57/59 Bill of Materials"
    else:
        suffix = "Bill of Materials"
        
    title = f"{customer} {sample_date} {suffix}"
    cell_title = ws['B1']
    cell_title.value = title
    cell_title.font = Font(name='Arial', bold=True, size=48)
    cell_title.alignment = Alignment(horizontal='center', vertical='center')

    cell_i1 = ws['I1']
    cell_i1.value = "Alimentación de materiales a tiempo\nMaterial feeding time\n送物料给生产时间"
    cell_i1.font = Font(name='Arial', bold=True, size=11)
    cell_i1.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    summary_headers = ["NO.", "SO#", "Customer", "", "Model", "Qty", "", "Date", "Remark"]
    ws.row_dimensions[2].height = 100
    for col_idx, val in enumerate(summary_headers, 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.value = val
        font_color = "FF0000" if col_idx == 9 else "000000"
        cell.font = Font(name='Arial', bold=True, size=24, color=font_color)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws.row_dimensions[3].height = 50

    for r_idx, row_data in enumerate(summary_rows, start=4):
        vals = [row_data["no"], row_data["so"], row_data["customer"],
                "", row_data["model"], row_data["qty"], "",
                row_data["date"], row_data["remark"]]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=r_idx, column=col, value=val)
            cell.font = Font(name='Arial', bold=True, size=24)
            if col == 9:
                cell.font = Font(name='Arial', bold=True, size=24, color='FF0000')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=(col == 3 or col == 7))
            cell.border = thin_border
        ws.row_dimensions[r_idx].height = 75

    detail_header_row = len(summary_rows) + 4
    detail_headers = [
        ("Número", "No.", "序号"),
        ("Orden de trabajo", "Work Order", "工单"),
        ("Número de pieza", "Part No.", "料号"),
        ("número de unidades", "Qty per WO", "需求数"),
        ("Cantidad demandada total", "Total Qty", "总需求数"),
        ("Unidad", "Unit", "单位"),
        ("Ubicación", "Location", "位置"),
        ("Cantidad en inventario", "Stock", "库存数"),
        ("Observación", "Remark", "备注"),
    ]
    for i, (es, en, zh) in enumerate(detail_headers):
        col = i + 1
        cell = ws.cell(row=detail_header_row, column=col)
        cell.value = f"{es}\n{en}\n{zh}"
        cell.font = Font(name='Arial', bold=True, size=18)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[detail_header_row].height = 88

    data_start_row = detail_header_row + 1
    current_pn = None
    for r_idx, row_data in enumerate(detail_rows, start=data_start_row):
        pn = row_data["pn"]
        is_first = (pn != current_pn)
        if is_first:
            current_pn = pn
            total_val = agg_dict.get(pn, 0.0)
        else:
            total_val = ""

        ws.cell(row=r_idx, column=1, value=row_data["no"])
        ws.cell(row=r_idx, column=2, value=row_data["wo"])
        ws.cell(row=r_idx, column=3, value=row_data["pn"])
        ws.cell(row=r_idx, column=4, value=row_data["req_qty"])
        ws.cell(row=r_idx, column=5, value=total_val)
        ws.cell(row=r_idx, column=6, value=row_data["unit"])
        cell_g = ws.cell(row=r_idx, column=7, value=row_data["bin"])
        cell_g.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.cell(row=r_idx, column=8, value=row_data["stock"])
        ws.cell(row=r_idx, column=9, value=row_data["remark"])

        ws.row_dimensions[r_idx].height = 70
        is_highlight = pn_groups.get(pn, {}).get("last_idx", -1) > pn_groups.get(pn, {}).get("first_idx", -1)
        for col in range(1, 10):
            cell = ws.cell(row=r_idx, column=col)
            cell.font = Font(name='Arial', bold=True, size=24)
            cell.alignment = Alignment(
                horizontal='center', vertical='center', wrap_text=(col == 7)
            )
            cell.border = thin_border
            if is_highlight:
                cell.fill = HIGHLIGHT_FILL
        if '\n' in (row_data["bin"] or ""):
            ws.row_dimensions[r_idx].height = None

    max_row = data_start_row + len(detail_rows) - 1
    ws.print_area = f"A1:{get_column_letter(9)}{max_row}"
    ws.oddFooter.center.text = "Page &[Page] of &[Pages]"
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = False
    ws.page_setup.orientation = 'portrait'
    ws.page_margins.left = 0
    ws.page_margins.right = 0
    ws.page_margins.top = 0
    ws.page_margins.bottom = 0
    ws.page_margins.header = 0
    ws.page_margins.footer = 0

    wb.save(filepath)
    print(f"[FORMATTED EXCEL] Saved -> {filepath}")

def generate_formatted_excel_52(records, bom_rows, agg_dict, inv_dict, station, prefix, filepath):
    customer = station.split('_')[0] if '_' in station else station
    wo_set = sorted({r["work_order"] for r in records}, key=work_order_sort_key)
    summary_rows = []
    for idx, wo in enumerate(wo_set, 1):
        rec = next(r for r in records if r["work_order"] == wo)
        raw_note = rec.get("ready_note", "")
        if raw_note.startswith("Please ready materials on "):
            remark = raw_note.replace("Please ready materials on ", "")
        else:
            remark = raw_note
        summary_rows.append({
            "no": idx,
            "so": wo,
            "customer": station,
            "model": rec["part_number"],
            "qty": rec["qty"],
            "date": rec["issue_date"],
            "remark": remark
        })

    detail_dict = {}
    for item in bom_rows:
        pn = item["partno_id"]
        if pn not in detail_dict:
            detail_dict[pn] = {"total_qty": 0.0}
        detail_dict[pn]["total_qty"] += item["qty"]

    detail_rows = []
    for idx, (pn, data) in enumerate(sorted(detail_dict.items()), 1):
        inv_info = inv_dict.get(pn, {"bin_id": "", "total_stock": 0.0})
        stock = inv_info["total_stock"]
        total_qty = data["total_qty"]
        detail_rows.append({
            "no": idx,
            "pn": pn,
            "bin": inv_info["bin_id"],
            "total_qty": total_qty,
            "stock": stock,
        })

    wb = Workbook()
    ws = wb.active
    ws.title = "Picking"
    col_widths = [20, 44, 38, 44, 50, 20, 20, 36]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    thin_border = Border(left=Side(style='thin', color='000000'), right=Side(style='thin', color='000000'),
                         top=Side(style='thin', color='000000'), bottom=Side(style='thin', color='000000'))
    for row in range(1, 3):
        for col in range(1, 9):
            ws.cell(row=row, column=col).border = thin_border

    cell_a1 = ws['A1']
    cell_a1.value = "Firma del obsequiador\nIssuer Signature\n执料人员签名"
    cell_a1.font = Font(name='Arial', bold=True, size=11)
    cell_a1.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 100

    ws.merge_cells('B1:G2')
    ws['A2'] = None
    for row in range(1, 3):
        for col in range(2, 8):
            ws.cell(row=row, column=col).border = thin_border
    ws.row_dimensions[2].height = 100
    raw_date = records[0]["issue_date"] if records else None
    sample_date = get_spanish_date_str(raw_date)
    title = f"{customer} {sample_date} 52 SMT Bill of Materials"
    cell_title = ws['B1']
    cell_title.value = title
    cell_title.font = Font(name='Arial', bold=True, size=48)
    cell_title.alignment = Alignment(horizontal='center', vertical='center')

    cell_h1 = ws['H1']
    cell_h1.value = "Alimentación de materiales a tiempo\nMaterial feeding time\n送物料给SMT时间"
    cell_h1.font = Font(name='Arial', bold=True, size=11)
    cell_h1.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    summary_headers = ["NO.", "SO#", "Customer", "Model", "B/T", "Qty", "Date", "Remark"]
    for col_idx, val in enumerate(summary_headers, 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.value = val
        font_color = "FF0000" if col_idx == 8 else "000000"
        cell.font = Font(name='Arial', bold=True, size=24, color=font_color)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws.row_dimensions[3].height = 57

    for r_idx, row_data in enumerate(summary_rows, start=4):
        vals = [row_data["no"], row_data["so"], row_data["customer"],
                row_data["model"], "", row_data["qty"],
                row_data["date"], row_data["remark"]]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=r_idx, column=col, value=val)
            cell.font = Font(name='Arial', bold=True, size=24)
            if col == 8:
                cell.font = Font(name='Arial', bold=True, size=24, color='FF0000')
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border
        ws.row_dimensions[r_idx].height = 75

    detail_header_row = len(summary_rows) + 4
    detail_headers_52 = [
        ("Número", "No.", "序号"),
        ("Número de pieza", "Part No.", "料号"),
        ("Ubicación", "Location", "位置"),
        ("Cantidad demandada total", "Total Qty", "总需求数"),
        ("实际发料数量", "Actual issued", "实际发料"),
        ("偏差", "Deviation", "偏差"),
        ("Inventario", "Stock", "库存总数"),
        ("Observación", "Remark", "备注")
    ]
    for i, (es, en, zh) in enumerate(detail_headers_52):
        col = i + 1
        cell = ws.cell(row=detail_header_row, column=col)
        cell.value = f"{es}\n{en}\n{zh}"
        cell.font = Font(name='Arial', bold=True, size=18)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[detail_header_row].height = 88

    data_start_row = detail_header_row + 1
    for r_idx, row_data in enumerate(detail_rows, start=data_start_row):
        ws.cell(row=r_idx, column=1, value=row_data["no"])
        ws.cell(row=r_idx, column=2, value=row_data["pn"])
        cell_c = ws.cell(row=r_idx, column=3, value=row_data["bin"])
        cell_c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.cell(row=r_idx, column=4, value=row_data["total_qty"])
        ws.cell(row=r_idx, column=5, value="")
        ws.cell(row=r_idx, column=6, value="")
        ws.cell(row=r_idx, column=7, value=row_data["stock"])
        ws.cell(row=r_idx, column=8, value="")

        ws.row_dimensions[r_idx].height = 70
        for col in range(1, 9):
            cell = ws.cell(row=r_idx, column=col)
            cell.font = Font(name='Arial', bold=True, size=24)
            cell.alignment = Alignment(
                horizontal='center', vertical='center', wrap_text=(col == 3)
            )
            cell.border = thin_border
        if '\n' in (row_data["bin"] or ""):
            ws.row_dimensions[r_idx].height = None

    max_row = data_start_row + len(detail_rows) - 1
    ws.print_area = f"A1:{get_column_letter(8)}{max_row}"
    ws.oddFooter.center.text = "Page &[Page] of &[Pages]"
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = False
    ws.page_setup.orientation = 'portrait'
    ws.page_margins.left = 0
    ws.page_margins.right = 0
    ws.page_margins.top = 0
    ws.page_margins.bottom = 0
    ws.page_margins.header = 0
    ws.page_margins.footer = 0

    wb.save(filepath)
    print(f"[FORMATTED EXCEL 52] Saved -> {filepath}")

# ===================== 原简单导出函数（备用） =====================
def export_final_picking_sheet(rows: List[Dict], filepath: str, sheet_name: str = "Picking"):
    rows_sorted = sorted(rows, key=lambda x: x["pn"])
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(["PN", "Bin_Location", "Req_Qty", "Total_Qty", "Stock_Qty"])
    for r in rows_sorted:
        ws.append([
            r["pn"],
            bin_to_multiline(r.get("bin_id", "")),
            r.get("req_qty", ""),
            r.get("total_qty", ""),
            r.get("stock_qty", "")
        ])
    style_worksheet(ws, [24, 30, 14, 14, 14])
    wb.save(filepath)
    print(f"[FINAL-EXCEL] Saved -> {filepath} | rows:{len(rows_sorted)}")

# ===================== 主流程与界面 =====================
def run_process(raw_clip: str, company_id: str, output_dir: str):
    global GLOBAL_COMPANY_ID
    GLOBAL_COMPANY_ID = company_id.strip()
    os.makedirs(output_dir, exist_ok=True)
    old_cwd = os.getcwd()
    os.chdir(output_dir)
    try:
        parsed_all = parse_clip_text(raw_clip)
        if not parsed_all:
            raise ValueError("No valid work orders found. Please check the clipboard format.")

        valid_records, dup_records = deduplicate_records(parsed_all)
        save_two_sheet_excel(valid_records, dup_records, "material_data.xlsx")
        save_zmf60a(valid_records, "ZMF60A.xls")
        zmf_path = os.path.abspath("ZMF60A.xls")
        pyperclip.copy(zmf_path)

        token = MES().login(USER_NAME, PASSWORD)
        if not token:
            raise RuntimeError("MES login failed. Please check the network or account settings.")

        station_group = group_by_station_and_prefix(valid_records)
        inv_dict = process_inventory(fetch_all_inventory(token))
        all_final_rows = []
        for station, group_data in station_group.items():
            for key, group in group_data.items():
                if not group:
                    continue
                filter_group = [rec for rec in group if "-P" not in rec["part_number"]]
                if not filter_group:
                    continue
                wo_set = sorted({r["work_order"] for r in filter_group}, key=work_order_sort_key)
                bom_rows = []
                for wo in wo_set:
                    bom_rows.extend(fetch_work_order_bom(token, wo))
                if not bom_rows:
                    print(f"[MAIN WARNING] station={station} key={key} got zero bom rows")
                    summary_file = f"{station}_{key}_{format_wo_name(wo_set)}_summary.xlsx"
                    save_summary_only(filter_group, station, key, wo_set, summary_file)
                    continue
                agg_dict = calc_agg_dict(bom_rows)
                out_file = f"{station}_{key}_{format_wo_name(wo_set)}.xlsx"
                if key == "prefix_52":
                    generate_formatted_excel_52(filter_group, bom_rows, agg_dict, inv_dict, station, key, out_file)
                else:
                    generate_formatted_excel(filter_group, bom_rows, agg_dict, inv_dict, station, key, out_file)
                for row in mark_last_row_total(bom_rows, agg_dict):
                    inv_info = inv_dict.get(row["partno_id"], {"bin_id":"", "total_stock":0.0})
                    all_final_rows.append({"pn": row["partno_id"], "bin_id": inv_info["bin_id"],
                                           "req_qty": row["qty"], "total_qty": row.get("total_qty"),
                                           "stock_qty": inv_info["total_stock"]})
        if all_final_rows:
            export_final_picking_sheet(all_final_rows, "bom_stock_result.xlsx", sheet_name="All_Picking")
        print(f"[COMPLETED] Output folder: {os.path.abspath(output_dir)}")
        return os.path.abspath(output_dir)
    finally:
        os.chdir(old_cwd)


SAMPLE_TEXT = """customer\tS/O Issue date\tSO#\tMODEL\tDeserition\tQTY\tPMC\t备注
PASCAL_VTX\t8/18a\t681000008937\t52-011385-001-PX00\tINDIRECT MATERIAL ASSY,U-A2S BOARD\t192\t\tPlease ready materials on 8/19 6am
	8/18a\t681000008936\t52-011385-001-RX00\tSMD ASSY,U-A2S\t192\t\tPlease ready materials on 8/19 6am"""
SAMPLE_IMAGE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_data.png")


def save_summary_only(records: List[Dict], station: str, key: str, wo_set: List[str], filepath: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary_Only"
    customer = station.split('_')[0] if '_' in station else station
    sample_date = get_spanish_date_str(records[0].get("issue_date") if records else None)
    suffix = {"prefix_50": "50", "prefix_80_85": "80&85", "prefix_55_57_59": "55/57/59"}.get(key, key)
    ws.merge_cells('A1:I1')
    title = ws['A1']
    title.value = f"{customer} {sample_date} {suffix} Bill of Materials (Summary Only)"
    title.font = Font(name='Arial', bold=True, size=20)
    title.alignment = Alignment(horizontal='center', vertical='center')
    headers = ["NO.", "SO#", "Customer", "Model", "Qty", "Date", "Remark"]
    for col_idx, value in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=value)
        cell.font = Font(name='Arial', bold=True, size=14)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = THIN_BORDER
    for idx, wo in enumerate(wo_set, 1):
        rec = next(row for row in records if row["work_order"] == wo)
        remark = rec.get("ready_note", "").replace("Please ready materials on ", "")
        for col_idx, value in enumerate([idx, wo, station, rec["part_number"], rec["qty"], rec["issue_date"], remark], 1):
            cell = ws.cell(row=3 + idx, column=col_idx, value=value)
            cell.font = Font(name='Arial', size=12)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = THIN_BORDER
    info_row = 5 + len(wo_set)
    ws.merge_cells(f'A{info_row}:I{info_row}')
    warning = ws[f'A{info_row}']
    warning.value = "WARNING: No BOM data available. Please check MES system."
    warning.font = Font(name='Arial', bold=True, size=14, color="FF0000")
    warning.alignment = Alignment(horizontal='center', vertical='center')
    for index, width in enumerate([10, 20, 25, 30, 10, 15, 30], 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    wb.save(filepath)
    print(f"[MAIN] Summary-only file created: {filepath}")


class GuiLogStream:
    def __init__(self, app):
        self.app = app
        self.original = app.original_stdout
        self.buffer = ""

    def write(self, text):
        self.original.write(text)
        self.original.flush()
        self.buffer += text
        while "\n" in self.buffer:
            line, self.buffer = self.buffer.split("\n", 1)
            if line.strip():
                self.app.root.after(0, self.app._append_log, line)

    def flush(self):
        self.original.flush()


class MaterialApp:
    def __init__(self, root):
        self.root = root
        root.title("MES Work Order Tool")
        root.geometry("980x700")
        root.minsize(760, 520)
        frame = ttk.Frame(root, padding=12)
        frame.pack(fill="both", expand=True)
        ttk.Label(frame, text="MES Work Order Tool", font=("Segoe UI", 16, "bold")).pack(anchor="w")
        config = ttk.Frame(frame)
        config.pack(fill="x", pady=(10, 6))
        ttk.Label(config, text="Company ID:").pack(side="left")
        self.company_var = tk.StringVar(value=GLOBAL_COMPANY_ID)
        company_box = ttk.Combobox(config, textvariable=self.company_var, values=("6501", "6801"), width=12)
        company_box.pack(side="left", padx=(6, 20))
        ttk.Button(config, text="Read Clipboard", command=self.load_clipboard).pack(side="left")
        ttk.Button(config, text="Open Sample Image", command=self.open_sample_image).pack(side="left", padx=6)
        self.start_button = ttk.Button(config, text="Start Processing", command=self.start)
        self.start_button.pack(side="left", padx=6)
        self.status_var = tk.StringVar(value="Copy the data, then start processing")
        ttk.Label(config, textvariable=self.status_var).pack(side="right")
        self.sample_image = Image.open(SAMPLE_IMAGE)
        image_frame = ttk.LabelFrame(frame, text="Clipboard Data Sample", padding=8)
        image_frame.pack(fill="both", expand=True, pady=(8, 0))
        self.image_frame = image_frame
        self.sample_label = ttk.Label(image_frame)
        self.sample_label.pack(fill="both", expand=True)
        image_frame.bind("<Configure>", self.resize_sample_image)
        ttk.Label(frame, text="Run log").pack(anchor="w", pady=(8, 2))
        self.log = tk.Text(frame, height=8, state="disabled", wrap="word")
        self.log.pack(fill="x")
        self.original_stdout = sys.stdout
        sys.stdout = GuiLogStream(self)
        self.root.after(100, self.resize_sample_image)

    def load_clipboard(self):
        try:
            value = pyperclip.paste()
            if not value.strip():
                raise ValueError("Clipboard is empty.")
            self.status_var.set("Clipboard loaded")
        except Exception as exc:
            messagebox.showerror("Read failed", str(exc))

    def open_sample_image(self):
        if not os.path.exists(SAMPLE_IMAGE):
            messagebox.showerror("Sample unavailable", SAMPLE_IMAGE)
            return
        os.startfile(SAMPLE_IMAGE)

    def resize_sample_image(self, _event=None):
        width = max(self.image_frame.winfo_width() - 16, 1)
        height = max(self.image_frame.winfo_height() - 16, 1)
        preview = self.sample_image.copy()
        preview.thumbnail((width, height), Image.Resampling.LANCZOS)
        self.sample_photo = ImageTk.PhotoImage(preview)
        self.sample_label.configure(image=self.sample_photo)

    def write_log(self, message):
        self.root.after(0, self._append_log, message)

    def _append_log(self, message):
        self.log.configure(state="normal")
        self.log.insert("end", message + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")

    def start(self):
        raw_clip = pyperclip.paste()
        company_id = self.company_var.get().strip()
        if not company_id.isdigit():
            messagebox.showerror("Invalid input", "Company ID must be numeric.")
            return
        if not raw_clip.strip():
            messagebox.showerror("Invalid input", "Clipboard data cannot be empty.")
            return
        date_dir = datetime.now().strftime("%Y-%m-%d_%H%M%S_%f")
        output_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), date_dir)
        self.start_button.configure(state="disabled")
        self.status_var.set("Processing, please wait...")
        threading.Thread(target=self.worker, args=(raw_clip, company_id, output_dir), daemon=True).start()

    def worker(self, raw_clip, company_id, output_dir):
        try:
            result_dir = run_process(raw_clip, company_id, output_dir)
            self.root.after(0, lambda: (self.status_var.set("Completed"), os.startfile(result_dir),
                                       messagebox.showinfo("Completed", f"Files saved and folder opened:\n{result_dir}")))
        except Exception as exc:
            self.write_log(str(exc))
            self.root.after(0, lambda: (self.status_var.set("Failed"), messagebox.showerror("Processing failed", str(exc))))
        finally:
            self.root.after(0, lambda: self.start_button.configure(state="normal"))


if __name__ == "__main__":
    MaterialApp(tk.Tk()).root.mainloop()